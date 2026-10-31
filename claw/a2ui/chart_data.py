"""Server-side data -> chart path for the `chart_from_data` tool.

The model only picks WHAT to chart (file, columns, aggregate, chart type,
title/labels). Every number - aggregates, percentages, period-over-period
deltas - is computed here with Decimal, then written as a Softnix `Chart`
card (.a2ui artifact). The model is told to quote these numbers verbatim.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from claw.a2ui import A2UI_VERSION, SOFTNIX_CATALOG_ID
from claw.a2ui.validate import (
    CHART_MAX_CATEGORIES,
    CHART_MAX_LABEL,
    CHART_MAX_PIE,
    CHART_MAX_POINTS,
    CHART_MAX_SERIES,
    CHART_TYPES,
)

AGGS = ("sum", "avg", "count", "min", "max")
GRAINS = ("day", "week", "month", "quarter", "year")
SORTS = ("auto", "x", "value_desc", "value_asc", "none")
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 200_000
SUFFIXES = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".json")


class ChartDataError(ValueError):
    """A problem the model can fix (bad column, too many categories...)."""


# ---- loading ---------------------------------------------------------------


def _rows_from_json(obj: Any) -> list[dict]:
    if isinstance(obj, dict):
        for key in ("rows", "data", "items", "records", "results"):
            if isinstance(obj.get(key), list):
                obj = obj[key]
                break
    if not isinstance(obj, list) or not all(isinstance(r, dict) for r in obj):
        raise ChartDataError("JSON must be a list of objects (or {\"rows\": [...]})")
    return obj


def load_rows(path: Path, sheet: str | None = None) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise ChartDataError(f"unsupported file type {suffix or '(none)'}; use {', '.join(SUFFIXES)}")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ChartDataError(f"file larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    if suffix == ".json":
        try:
            return _rows_from_json(json.loads(path.read_text(encoding="utf-8-sig")))[:MAX_ROWS]
        except json.JSONDecodeError as exc:
            raise ChartDataError(f"invalid JSON ({exc.msg} at line {exc.lineno})") from exc
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet:
                if sheet not in wb.sheetnames:
                    raise ChartDataError(f"sheet {sheet!r} not found; sheets: {wb.sheetnames}")
                ws = wb[sheet]
            else:
                ws = wb.worksheets[0]
            it = ws.iter_rows(values_only=True)
            header: list[str] | None = None
            out: list[dict] = []
            for row in it:
                if header is None:
                    if row and any(v not in (None, "") for v in row):
                        header = [str(v).strip() if v is not None else f"col{i + 1}" for i, v in enumerate(row)]
                    continue
                if not any(v not in (None, "") for v in row):
                    continue
                out.append({h: (row[i] if i < len(row) else None) for i, h in enumerate(header)})
                if len(out) >= MAX_ROWS:
                    break
            return out
        finally:
            wb.close()
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    out = []
    for row in reader:
        out.append({(k or "").strip(): v for k, v in row.items() if k is not None})
        if len(out) >= MAX_ROWS:
            break
    return out


# ---- parsing ---------------------------------------------------------------

_NUM_STRIP = re.compile(r"[,\s฿$€£¥₩₹]|THB|USD|บาท", re.I)


def parse_number(v: Any) -> Decimal | None:
    """Decimal from a cell: 1,234.50 / ฿1,234 / (12) / 15% / '' -> None. Raises on junk."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, Decimal)):
        return Decimal(v)
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return None
        return Decimal(repr(v))
    s = str(v).strip()
    if s in ("", "-", "–", "—", "n/a", "N/A", "NA", "null", "None"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = _NUM_STRIP.sub("", s.strip("()")).rstrip("%")
    try:
        d = Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(v) from exc
    if not d.is_finite():
        raise ValueError(v)
    return -d if neg else d


_DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m", "%Y/%m", "%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S")


def parse_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("Z", "")[:19]
    for f in _DATE_FORMATS:
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s).date()
    except ValueError:
        return None


_TH_MONTHS = ("ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค.")
_EN_MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


def bucket(d: date, grain: str, lang: str) -> tuple[date, str]:
    """(sort key, label) of the period containing d."""
    th = lang == "th"
    if grain == "day":
        return d, d.strftime("%d/%m/%Y" if th else "%Y-%m-%d")
    if grain == "week":
        start = d - timedelta(days=d.weekday())
        y, w, _ = d.isocalendar()
        return start, (f"สัปดาห์ {w}/{y}" if th else f"{y}-W{w:02d}")
    if grain == "month":
        names = _TH_MONTHS if th else _EN_MONTHS
        return date(d.year, d.month, 1), f"{names[d.month - 1]} {d.year}"
    if grain == "quarter":
        q = (d.month - 1) // 3 + 1
        return date(d.year, 3 * q - 2, 1), (f"ไตรมาส {q}/{d.year}" if th else f"Q{q} {d.year}")
    return date(d.year, 1, 1), str(d.year)


# ---- spec + aggregation ----------------------------------------------------


@dataclass
class ChartSpec:
    chart_type: str
    x: str
    y: list[str] = field(default_factory=list)
    agg: str = "sum"
    series_by: str | None = None
    x_grain: str | None = None
    where: dict[str, Any] = field(default_factory=dict)
    sort: str = "auto"
    limit: int | None = None
    delta: str = "none"
    good_direction: str = "up"
    title: str = ""
    x_label: str = ""
    y_label: str = ""
    unit: str = ""
    decimals: int | None = None
    series_labels: dict[str, str] = field(default_factory=dict)
    lang: str = "th"
    show_table: bool = False


@dataclass
class ChartResult:
    categories: list[str]
    series: list[tuple[str, list[Decimal | None]]]
    deltas: list[list[dict | None]] | None
    rows_used: int
    rows_total: int
    skipped: int
    decimals: int
    merged_other: int = 0
    notes: list[str] = field(default_factory=list)


def _col(columns: list[str], name: str, what: str) -> str:
    if name in columns:
        return name
    low = {c.strip().lower(): c for c in columns}
    hit = low.get(name.strip().lower())
    if hit is None:
        raise ChartDataError(f"{what}: column {name!r} not found; columns are {columns[:40]}")
    return hit


def _q(d: Decimal, places: int) -> Decimal:
    return d.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)


def pct_change(cur: Decimal | None, prev: Decimal | None) -> dict | None:
    if cur is None or prev is None:
        return None
    if prev == 0:
        return {"pct": None, "direction": "flat" if cur == 0 else "up" if cur > 0 else "down"}
    pct = _q((cur - prev) / abs(prev) * 100, 1)
    return {"pct": float(pct), "direction": "up" if pct > 0 else "down" if pct < 0 else "flat"}


def validate_spec(spec: ChartSpec) -> None:
    if spec.chart_type not in CHART_TYPES:
        raise ChartDataError(f"chart_type: one of {', '.join(CHART_TYPES)}")
    if spec.agg not in AGGS:
        raise ChartDataError(f"agg: one of {', '.join(AGGS)}")
    if spec.x_grain is not None and spec.x_grain not in GRAINS:
        raise ChartDataError(f"x_grain: one of {', '.join(GRAINS)}")
    if spec.sort not in SORTS:
        raise ChartDataError(f"sort: one of {', '.join(SORTS)}")
    if spec.delta not in ("none", "previous"):
        raise ChartDataError("delta: none|previous")
    if spec.series_by and len(spec.y) > 1:
        raise ChartDataError("series_by works with exactly one y column")
    if len(spec.y) > CHART_MAX_SERIES:
        raise ChartDataError(f"y: at most {CHART_MAX_SERIES} columns")
    if spec.chart_type in ("pie", "donut") and (len(spec.y) > 1 or spec.series_by):
        raise ChartDataError(f"{spec.chart_type}: one y column and no series_by")
    if spec.limit is not None and not (1 <= spec.limit <= CHART_MAX_CATEGORIES):
        raise ChartDataError(f"limit: 1-{CHART_MAX_CATEGORIES}")
    for name, v, cap in (("title", spec.title, 200), ("x_label", spec.x_label, CHART_MAX_LABEL), ("y_label", spec.y_label, CHART_MAX_LABEL), ("unit", spec.unit, 20)):
        if len(v or "") > cap:
            raise ChartDataError(f"{name}: max {cap} characters")


def _looks_numeric(rows: list[dict], col: str) -> bool:
    seen = 0
    for row in rows[:50]:
        try:
            if parse_number(row.get(col)) is not None:
                seen += 1
        except ValueError:
            return False
    return seen > 0


def aggregate(rows: list[dict], spec: ChartSpec) -> ChartResult:
    validate_spec(spec)
    if not rows:
        raise ChartDataError("the file has no data rows")
    columns = list(rows[0].keys())
    notes: list[str] = []
    if spec.agg != "count" and not spec.y:
        numeric = [
            c for c in columns if c.lower() not in (spec.x.lower(), (spec.series_by or "").lower()) and _looks_numeric(rows, c)
        ]
        if len(numeric) == 1:  # unambiguous: the only number column
            spec.y = numeric
            notes.append(f"y was not given; used the only numeric column {numeric[0]!r}.")
        else:
            raise ChartDataError(
                f"y is missing: agg='{spec.agg}' needs the number column(s) to aggregate, e.g. y='{numeric[0] if numeric else 'amount'}'"
                f" (numeric columns: {numeric[:20]}). To count rows per {spec.x!r} instead, set agg='count'."
            )
    x = _col(columns, spec.x, "x")
    ys = [_col(columns, y, "y") for y in spec.y]
    sb = _col(columns, spec.series_by, "series_by") if spec.series_by else None
    where = {_col(columns, k, "where"): v for k, v in (spec.where or {}).items()}
    other_label = "อื่น ๆ" if spec.lang == "th" else "Other"
    count_label = "จำนวน" if spec.lang == "th" else "Count"
    blank = "(ว่าง)" if spec.lang == "th" else "(blank)"

    # (xkey, series) -> [sum, count, min, max]
    acc: dict[tuple[Any, str], list] = {}
    x_labels: dict[Any, str] = {}
    x_order: list[Any] = []
    s_order: list[str] = []
    skipped = used = 0
    for row in rows:
        if any(str(row.get(k, "")).strip() != str(v).strip() for k, v in where.items()):
            continue
        raw_x = row.get(x)
        if spec.x_grain:
            d = parse_date(raw_x)
            if d is None:
                skipped += 1
                continue
            xkey, xlabel = bucket(d, spec.x_grain, spec.lang)
        else:
            xlabel = str(raw_x).strip() if raw_x not in (None, "") else blank
            if isinstance(raw_x, float) and raw_x.is_integer():
                xlabel = str(int(raw_x))
            xkey = xlabel
        pairs: list[tuple[str, Decimal | None]] = []
        try:
            if spec.agg == "count" and not ys:
                pairs = [((str(row.get(sb, "")).strip() or blank) if sb else count_label, Decimal(1))]
            else:
                for y in ys:
                    if spec.agg == "count":  # count non-empty cells; no number parsing
                        cell = row.get(y)
                        val = None if cell is None or str(cell).strip() == "" else Decimal(1)
                    else:
                        val = parse_number(row.get(y))
                    label = (str(row.get(sb, "")).strip() or blank) if sb else (count_label if spec.agg == "count" else y)
                    pairs.append((label, Decimal(1) if spec.agg == "count" and val is not None else val))
        except ValueError:
            skipped += 1
            continue
        if all(v is None for _, v in pairs):
            skipped += 1
            continue
        used += 1
        if xkey not in x_labels:
            x_labels[xkey] = xlabel[:CHART_MAX_LABEL]
            x_order.append(xkey)
        for label, v in pairs:
            if v is None:
                continue
            if label not in s_order:
                s_order.append(label)
            a = acc.setdefault((xkey, label), [Decimal(0), 0, None, None])
            a[0] += v
            a[1] += 1
            a[2] = v if a[2] is None or v < a[2] else a[2]
            a[3] = v if a[3] is None or v > a[3] else a[3]
    if not used:
        raise ChartDataError(f"no usable rows ({skipped} skipped: missing/non-numeric values{' or unparseable dates' if spec.x_grain else ''})")

    def value(a: list | None) -> Decimal | None:
        if a is None:
            return None
        if spec.agg == "sum":
            return a[0]
        if spec.agg == "count":
            return Decimal(a[1]) if spec.y else a[0]
        if spec.agg == "avg":
            return a[0] / a[1]
        return a[2] if spec.agg == "min" else a[3]

    table = {s: [value(acc.get((k, s))) for k in x_order] for s in s_order}
    totals = {k: sum((table[s][i] or Decimal(0)) for s in s_order) for i, k in enumerate(x_order)}

    # order categories
    sort = spec.sort
    if sort == "auto":
        sort = "x" if spec.x_grain else "value_desc" if spec.chart_type in ("pie", "donut", "hbar", "progress") else "none"
    idx = list(range(len(x_order)))
    if sort == "x":
        idx.sort(key=lambda i: (x_order[i] if spec.x_grain else str(x_order[i])))
    elif sort in ("value_desc", "value_asc"):
        idx.sort(key=lambda i: totals[x_order[i]], reverse=(sort == "value_desc"))
    merged = 0
    cap = CHART_MAX_PIE if spec.chart_type in ("pie", "donut") else CHART_MAX_CATEGORIES
    if spec.limit:
        cap = min(cap, spec.limit)
    keep, rest = idx[:cap], idx[cap:]
    if rest:
        if spec.chart_type in ("pie", "donut") and spec.agg in ("sum", "count"):
            keep, rest = idx[: cap - 1], idx[cap - 1 :]
            merged = len(rest)
        elif not spec.limit:
            raise ChartDataError(
                f"{len(idx)} categories is too many (max {cap}); set limit (top N), x_grain (day/week/month/...), or where"
            )
    categories = [x_labels[x_order[i]] for i in keep]
    series_vals = {s: [table[s][i] for i in keep] for s in s_order}
    if merged:
        categories.append(other_label)
        for s in s_order:
            series_vals[s].append(sum((table[s][i] or Decimal(0)) for i in rest))

    # cap series (series_by can explode)
    names = list(s_order)
    if len(names) > CHART_MAX_SERIES:
        if spec.agg not in ("sum", "count"):
            raise ChartDataError(f"series_by gives {len(names)} series (max {CHART_MAX_SERIES}); filter with where")
        by_total = sorted(names, key=lambda s: sum(v or 0 for v in series_vals[s]), reverse=True)
        top, low = by_total[: CHART_MAX_SERIES - 1], by_total[CHART_MAX_SERIES - 1 :]
        series_vals[other_label] = [
            sum((series_vals[s][i] or Decimal(0)) for s in low) for i in range(len(categories))
        ]
        names = top + [other_label]
    if len(categories) * len(names) > CHART_MAX_POINTS:
        raise ChartDataError(f"{len(categories)} x {len(names)} points is too many (max {CHART_MAX_POINTS})")

    if spec.decimals is not None:
        decimals = spec.decimals
    elif spec.agg == "count" or all(v is None or v == v.to_integral_value() for s in names for v in series_vals[s]):
        decimals = 0  # whole numbers (counts, token totals) never show ".00"
    else:
        decimals = 2
    series = [
        (spec.series_labels.get(s, s)[:CHART_MAX_LABEL], [None if v is None else _q(v, decimals) for v in series_vals[s]])
        for s in names
    ]
    deltas = None
    if spec.delta == "previous":
        deltas = [[None] + [pct_change(vals[i], vals[i - 1]) for i in range(1, len(vals))] for _, vals in series]
    return ChartResult(categories, series, deltas, used, len(rows), skipped, decimals, merged, notes)


# ---- card ------------------------------------------------------------------

_AGG_TXT = {
    "th": {"sum": "ผลรวม", "avg": "ค่าเฉลี่ย", "count": "จำนวนแถว", "min": "ค่าต่ำสุด", "max": "ค่าสูงสุด"},
    "en": {"sum": "sum", "avg": "average", "count": "count", "min": "minimum", "max": "maximum"},
}


def caption(spec: ChartSpec, res: ChartResult, source: str) -> str:
    th = spec.lang == "th"
    agg = _AGG_TXT["th" if th else "en"][spec.agg]
    what = ", ".join(spec.y) if spec.y else ("แถว" if th else "rows")
    grain = f" ({spec.x_grain})" if spec.x_grain else ""
    if th:
        txt = f"ที่มา: {source} · {agg}ของ {what} ตาม {spec.x}{grain} · ใช้ {res.rows_used:,} จาก {res.rows_total:,} แถว"
        if res.skipped:
            txt += f" (ข้าม {res.skipped:,} แถวที่ค่าว่าง/ไม่ใช่ตัวเลข)"
        if res.deltas is not None:
            txt += " · ▲▼ เทียบกับช่วงก่อนหน้า"
        txt += " · คำนวณโดยระบบ ไม่ใช่ AI"
    else:
        txt = f"Source: {source} · {agg} of {what} by {spec.x}{grain} · {res.rows_used:,} of {res.rows_total:,} rows used"
        if res.skipped:
            txt += f" ({res.skipped:,} skipped: empty/non-numeric)"
        if res.deltas is not None:
            txt += " · ▲▼ vs previous period"
        txt += " · computed by the system, not the AI"
    return txt


def build_messages(spec: ChartSpec, res: ChartResult, source: str, surface_id: str = "chart") -> list[dict]:
    chart: dict[str, Any] = {
        "id": "chart",
        "component": "Chart",
        "chartType": spec.chart_type,
        "categories": res.categories,
        "series": [
            {"label": label, "values": [None if v is None else float(v) for v in vals]}
            | ({"deltas": res.deltas[j]} if res.deltas is not None else {})
            for j, (label, vals) in enumerate(res.series)
        ],
        "decimals": res.decimals,
        "goodDirection": "down" if spec.good_direction == "down" else "up",
    }
    for key, val in (("unit", spec.unit), ("xLabel", spec.x_label), ("yLabel", spec.y_label)):
        if val:
            chart[key] = val
    if spec.show_table:
        chart["showTable"] = True
    comps: list[dict] = [
        {"id": "root", "component": "Card", "child": "col"},
        {"id": "col", "component": "Column", "children": ["title", "cap", "chart"]},
        {"id": "title", "component": "Text", "variant": "h3", "text": (spec.title or source)[:200]},
        {"id": "cap", "component": "Text", "variant": "caption", "text": caption(spec, res, source)},
        chart,
    ]
    return [
        {"version": A2UI_VERSION, "createSurface": {"surfaceId": surface_id, "catalogId": SOFTNIX_CATALOG_ID}},
        {"version": A2UI_VERSION, "updateComponents": {"surfaceId": surface_id, "components": comps}},
    ]


def _plain(d: Decimal | None) -> str:
    if d is None:
        return "-"
    return f"{d:,f}"


def summary_for_model(spec: ChartSpec, res: ChartResult, artifact: str) -> str:
    head = ["category"] + [label for label, _ in res.series]
    if res.deltas is not None:
        head += [f"Δ {label}" for label, _ in res.series]
    lines = [
        f"OK. Chart card written to {artifact} and shown to the user (chart + table toggle are already on screen).",
        f"{spec.chart_type} of {spec.agg}({', '.join(spec.y) or 'rows'}) by {spec.x}"
        + (f" per {spec.x_grain}" if spec.x_grain else "")
        + f"; {res.rows_used} of {res.rows_total} rows used, {res.skipped} skipped."
        + (f" {res.merged_other} small categories merged into the last one." if res.merged_other else "")
        + "".join(f" Note: {n}" for n in res.notes),
        "All numbers below were computed by the tool. Do NOT recompute, round differently, total up or invent numbers;",
        "quote them exactly if you mention any. Reply with a short summary (2-4 sentences) written entirely in the",
        "language of the user's latest message (no words from other languages).",
        " | ".join(head),
    ]
    for i, cat in enumerate(res.categories):
        cells = [cat] + [_plain(vals[i]) for _, vals in res.series]
        if res.deltas is not None:
            for j in range(len(res.series)):
                d = res.deltas[j][i]
                cells.append("-" if d is None else (f"{d['direction']} {d['pct']:+.1f}%" if d["pct"] is not None else d["direction"]))
        lines.append(" | ".join(cells))
    return "\n".join(lines)
