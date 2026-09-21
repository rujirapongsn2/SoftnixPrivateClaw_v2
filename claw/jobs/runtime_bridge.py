"""Worker boundary for the existing, policy-enforcing normal chat runtime."""
from dataclasses import asdict
from pathlib import Path
import hashlib
import uuid

from claw.jobs.contracts import JobPolicy, StepOutcome
from claw.jobs.provider import current_execution


class Promoted(BaseException):
    def __init__(self, job_id, turn_id):
        self.job_id, self.turn_id = job_id, turn_id


class Handoff(BaseException):
    """Exit nested tool handlers without presenting scheduling as a tool error."""
    def __init__(self, outcome):
        self.outcome = outcome


def pair_checkpoint_messages(messages):
    """Close unsent members of a tool batch without re-executing completed calls."""
    result = list(messages)
    answered = {m.get('tool_call_id') for m in result if m.get('role') == 'tool'}
    # A persisted checkpoint can end partway through a bundled tool response.
    for message in messages:
        for call in message.get('tool_calls', []):
            if call.get('id') not in answered:
                result.append({'role': 'tool', 'tool_call_id': call['id'],
                               'name': call.get('function', {}).get('name', ''),
                               'content': 'Deferred at checkpoint; this call has no confirmed result. Inspect current state before continuing.'})
    return result


class CheckpointStore:
    """Context-local facade; concurrent turns never replace a runtime's store."""
    def __init__(self, legacy):
        self.legacy = legacy

    def __getattr__(self, name):
        return getattr(self.legacy, name)

    async def update(self, job_id, **changes):
        ctx = current_execution.get()
        if ctx is None or job_id != ctx.lease.job_id:
            return await self.legacy.update(job_id, **changes)
        value = {**ctx.state.get('runtime', {}), **changes}
        await ctx.checkpoint({**ctx.state, 'runtime': value})
        return value

    async def finish(self, job_id, status, **changes):
        ctx = current_execution.get()
        if ctx is None or job_id != ctx.lease.job_id:
            return await self.legacy.finish(job_id, status, **changes)
        await self.update(job_id, status=status, **changes)
        raise Handoff(StepOutcome('paused', checkpoint=ctx.state, reason='runtime_' + status))


async def admit(runtime, user_id, session_id, content, locale, media, model, permission_mode, kind='artifact', *, adoption=None, job_id=None):
    session = await runtime.sessions.get(session_id)
    if session is None or session.user_id != user_id:
        raise PermissionError('session ownership mismatch')
    configured = runtime.settings.team_work
    saved = await runtime.artifact_jobs.resource_policy()
    policy = JobPolicy(max_tokens=saved.get('max_job_tokens', configured.max_job_tokens),
                       max_seconds=saved.get('max_job_seconds', configured.max_job_seconds),
                       max_recoveries=saved.get('max_step_recoveries', configured.max_step_recoveries))
    job_id = job_id or uuid.uuid4().hex
    workspace = runtime.get_agent(user_id).workspace.resolve()
    media = [str(Path(p).resolve().relative_to(workspace)) if Path(p).is_absolute() else p for p in media or []]
    inputs = dict(content=content, locale=locale, media=list(media or []), model=model,
                  permission_mode=permission_mode, session_id=session_id)
    from claw.core.context import build_user_content
    _, storage_text = build_user_content(content, media, runtime.get_agent(user_id).workspace)
    await runtime.durable_jobs.submit(job_id, user_id, session_id, 'privateclaw', [dict(
        id='task', executor='privateclaw.turn', effect='local', replay_safe=True,
        inputs=inputs, acceptance={'kind': kind})], policy=policy.model_dump(), locale=locale,
        user_message=storage_text if adoption is None else None, adoption=adoption)
    return job_id


async def promote_if_planned(runtime, user_id, session_id, content, locale, media, model,
                             permission_mode, state, turn_id):
    """Promote only after a completed plan tool and a supported acceptance contract."""
    from claw.jobs.provider import foreground_accounting
    from claw.core.runtime import _is_artifact_task
    from claw.jobs.research import is_multistep_research
    tape = foreground_accounting.get()
    if tape is None or not tape.calls:
        return
    messages = state['messages']
    if not messages or messages[-1].get('name') != 'update_plan':
        return
    if messages[-1].get('content', '').startswith('Error'):
        return
    session = await runtime.sessions.get(session_id)
    plan = session.plan or {}
    if len([s for s in plan.get('steps', []) if s.get('status') != 'done']) < 2:
        return
    # Other task families stay on their existing route until a corresponding
    # acceptance validator is wired. A model's plan alone is not proof of success.
    goal = str(plan.get('goal', ''))
    kind = 'artifact' if _is_artifact_task(goal, media) else 'research' if is_multistep_research(goal) else None
    if kind is None:
        return
    job_id = uuid.uuid4().hex
    adoption = tape.snapshot()
    adoption['checkpoint'] = {**tape.state, 'runtime': {
        'checkpoint_messages': messages, 'tool_results': state['tool_results'],
        'written': state['written'], 'pending_call': None,
        'user_message_persisted': True,
    }}
    await admit(runtime, user_id, session_id, content, locale, media, model,
                permission_mode, kind=kind, adoption=adoption, job_id=job_id)
    raise Promoted(job_id, turn_id)


def file_evidence(root: Path, paths):
    refs = []
    for name in dict.fromkeys(paths):
        path = (root / name).resolve()
        if not path.is_relative_to(root.resolve()) or not path.is_file():
            raise ValueError('missing or unsafe artifact')
        if path.suffix.lower() == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True)
            try:
                if not book.sheetnames or any(c.data_type == 'e' for s in book for row in s for c in row):
                    raise ValueError('invalid workbook')
            finally:
                book.close()
        refs.append({'path': name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    return {'files': refs, 'validation_scope': 'file integrity and workbook structure only'}


def sources_unchanged(root, state):
    try:
        for name, value in state.get('sources', {}).items():
            path = (root / name).resolve()
            if not path.is_relative_to(root.resolve()) or hashlib.sha256(path.read_bytes()).hexdigest() != value['sha256']:
                return False
        return True
    except (OSError, KeyError):
        return False


async def normal_execute(runtime, ctx):
    args = ctx.lease.spec.inputs
    if ctx.state.get('candidate'):
        return StepOutcome(**ctx.state['candidate'])
    pending = ctx.state.get('pending_tool') or ctx.state.get('runtime', {}).get('pending_call')
    # Unknown tools are unsafe to replay. A command returning no response may
    # already have written to an external system, even when named "exec".
    reads = {'read_file', 'list_dir', 'read_docx', 'read_pdf', 'read_skill', 'web_search', 'web_fetch'}
    if pending and pending.get('name') not in reads:
        return StepOutcome('paused', checkpoint=ctx.state, reason='uncertain_effect')
    initial = dict(id=ctx.lease.job_id, turn_id=ctx.lease.job_id[:12],
                   user_id=ctx.lease.owner_id, session_id=args['session_id'], status='running',
                   user_message_persisted=True, content=args['content'], segment=1,
                   budget={'tokens': 5_000_000, 'seconds': 21_600})
    from claw.jobs.models import Job
    async with ctx.store.factory() as db:
        job = await db.get(Job, ctx.lease.job_id)
        initial['budget'] = {'tokens': job.policy['max_tokens'], 'seconds': job.policy['max_seconds']}
    value = {**initial, **ctx.state.get('runtime', {})}
    value['checkpoint_messages'] = pair_checkpoint_messages(value.get('checkpoint_messages', []))
    await ctx.checkpoint({**ctx.state, 'runtime': value})
    instruction = args['content']
    if ctx.lease.spec.acceptance.get('kind') == 'research':
        instruction += ('\n\n[Research delivery contract: retrieve at least two relevant source pages with web_fetch. '
                        'Cite only URLs actually retrieved. Deliver the answer in chat; a file is not required. '
                        'Source/citation checks do not establish factual completeness. If evidence is missing, '
                        'report the limitation rather than claiming complete verification.]')
    try:
        await runtime._process_turn(ctx.lease.owner_id, args['session_id'], instruction,
            locale=args['locale'], media=args['media'], model=args['model'],
            permission_mode=args['permission_mode'], artifact_job=value)
    except Handoff as signal:
        return signal.outcome
    return StepOutcome('paused', checkpoint=ctx.state, reason='runtime_rejected')


async def handoff_outcome(outcome, root: Path | None = None):
    ctx = current_execution.get()
    if ctx is None:
        return
    if outcome.blocked_reason:
        raise Handoff(StepOutcome('dependency', checkpoint=ctx.state,
                                 dependency='sandbox', reason='dependency_unavailable'))
    if outcome.usage_limit_reached:
        raise Handoff(StepOutcome('paused', checkpoint=ctx.state, reason='resource_limit'))
    if outcome.timed_out or outcome.reached_max_iterations:
        evidence = progress_evidence(root, ctx.state) if root is not None else coverage_evidence(ctx.state)
        raise Handoff(StepOutcome('yielded', checkpoint=ctx.state, evidence=evidence, reason='slice_timeout'))


def coverage_evidence(state):
    return {'coverage': {key: {'sha256': value['sha256'], 'length': value['length'],
                               'ranges': value['ranges']}
                         for key, value in state.get('sources', {}).items()},
            'research': {key: value['sha256'] for key, value in state.get('research_sources', {}).items()
                         if value.get('origin') == 'web_fetch' and 200 <= value.get('status', 0) < 300
                         and hashlib.sha256(value.get('text', '').encode()).hexdigest() == value.get('sha256')}}


def progress_evidence(root: Path, state: dict) -> dict:
    """Return deterministic progress evidence without claiming completion.

    File count or model prose is not evidence. Each recorded file must still be
    inside the workspace and match its current hash; source and research entries
    retain their existing coverage/provenance checks.
    """
    evidence = coverage_evidence(state)
    runtime = state.get('runtime', {})
    written = list(dict.fromkeys([
        *state.get('written', []),
        *runtime.get('written', []),
    ]))
    files = []
    for name in written:
        try:
            files.extend(file_evidence(root, [name])['files'])
        except (OSError, ValueError, KeyError, TypeError):
            continue
    evidence['files'] = files
    return evidence


def validate_progress(root: Path, state: dict, evidence: dict) -> bool:
    """Independently re-check evidence used to grant another execution slice."""
    expected = progress_evidence(root, state)
    return bool(expected['coverage'] or expected['research'] or expected['files']) and evidence == expected


def coverage_complete(state, required):
    from claw.jobs.validators import source_coverage
    manifest = coverage_evidence(state)['coverage']
    docs = [name for name in required if name.lower().endswith('.docx')]
    return all(name in manifest and source_coverage({name: manifest[name]},
        {name: {'sha256': manifest[name]['sha256'], 'length': manifest[name]['length']}}) for name in docs)


async def handoff_delivery(root, final, artifacts):
    ctx = current_execution.get()
    if ctx is None:
        return
    if not sources_unchanged(root, ctx.state):
        raise Handoff(StepOutcome('paused', checkpoint=ctx.state, reason='source_changed'))
    if not coverage_complete(ctx.state, ctx.lease.spec.inputs.get('media', [])):
        raise Handoff(StepOutcome('paused', checkpoint=ctx.state, reason='source_coverage_incomplete'))
    if ctx.lease.spec.acceptance.get('kind') == 'research':
        from claw.jobs.research import research_evidence
        evidence = research_evidence(ctx.state, final)
        if evidence is None:
            raise Handoff(StepOutcome('paused', checkpoint=ctx.state, reason='validation_failed'))
        result = StepOutcome('completed', checkpoint=ctx.state, evidence=evidence,
            delivery={'content': final, 'meta': {'validation_scope': evidence['validation_scope']}})
        await ctx.checkpoint({**ctx.state, 'candidate': asdict(result)})
        raise Handoff(result)
    try:
        evidence = file_evidence(root, artifacts)
    except (ValueError, OSError):
        raise Handoff(StepOutcome('paused', checkpoint=ctx.state, reason='validation_failed')) from None
    if not evidence['files']:
        raise Handoff(StepOutcome('paused', checkpoint=ctx.state, reason='missing_deliverable'))
    result = StepOutcome('completed', checkpoint=ctx.state, evidence=evidence,
                         delivery={'content': final, 'meta': {'artifacts': artifacts,
                                   'validation_scope': evidence['validation_scope']}})
    await ctx.checkpoint({**ctx.state, 'candidate': asdict(result)})
    raise Handoff(result)
