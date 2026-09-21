"""Deterministic evidence checks. No validator accepts model self-attestation."""
import hashlib
from pathlib import Path


def workspace_file(root: Path, ref: dict) -> Path:
    root = root.resolve()
    path = (root / ref['path']).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ValueError('output is outside workspace or missing')
    if hashlib.sha256(path.read_bytes()).hexdigest() != ref['sha256']:
        raise ValueError('output hash does not match checkpoint')
    return path


def source_coverage(manifest: dict, required: dict) -> bool:
    """Require continuous coverage of each versioned input, not just a page count."""
    if not required or set(manifest) != set(required):
        return False
    for key, source in required.items():
        item = manifest[key]
        if item.get('sha256') != source['sha256'] or source['length'] < 0:
            return False
        end = 0
        for start, stop in sorted(item.get('ranges', [])):
            if start > end or stop < start or start < 0 or stop > source['length']:
                return False
            end = max(end, stop)
        if end != source['length']:
            return False
    return True


def validate_workbook(root: Path, ref: dict, expected_sheets: dict[str, list[str]]) -> bool:
    """Check structure and explicit error cells. Does not claim semantic/formula correctness."""
    from openpyxl import load_workbook
    if not expected_sheets:
        return False
    path = workspace_file(root, ref)
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if set(book.sheetnames) != set(expected_sheets):
            return False
        for name, headers in expected_sheets.items():
            sheet = book[name]
            if list(next(sheet.values, ()))[:len(headers)] != headers:
                return False
            if any(cell.data_type == 'e' for row in sheet for cell in row):
                return False
        return True
    finally:
        book.close()


def validate_research(rows: list[dict], required_questions: set[str], sources: dict) -> bool:
    """Coverage and citation provenance only; truth assessment requires a separate reviewer."""
    if not required_questions or {r.get('question_id') for r in rows} != required_questions:
        return False
    return all(row.get('answer') and row.get('citations') and
               all(c in sources and sources[c].get('sha256') for c in row['citations']) for row in rows)


def validate_receipt(receipt: dict, expected_key: str) -> bool:
    """Caller must retrieve the receipt using a trusted connector, not model text."""
    return bool(expected_key and receipt.get('idempotency_key') == expected_key
                and receipt.get('remote_id') and receipt.get('status') == 'confirmed')
