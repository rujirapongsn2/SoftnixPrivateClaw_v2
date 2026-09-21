"""Deterministic workbook generation from a schema-checked intermediate file.

Data cells are literal text/numbers, never implicit spreadsheet formulas. This
runs trusted application code, not model-authored Python on the API host.
"""
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
import asyncio
import sys
from pathlib import Path
import re
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from claw.tools.base import Tool

VERSION = 'workbook-v1'


def generate(root, source, output):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    root = Path(root).resolve()
    src, dst = (root / source).resolve(), (root / output).resolve()
    if not src.is_relative_to(root) or not dst.is_relative_to(root) or dst.suffix.lower() != '.xlsx' or src == dst:
        raise ValueError('paths must stay inside workspace; output must be .xlsx')
    if src.stat().st_size > 20_000_000:
        raise ValueError('intermediate data exceeds 20 MB')
    raw = src.read_bytes()
    data = json.loads(raw)
    sheets = data.get('sheets')
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 50:
        raise ValueError('sheets must contain 1–50 sheets')
    names, rows_total = set(), 0
    book = Workbook()
    book.remove(book.active)
    for spec in sheets:
        name, columns, rows = spec.get('name'), spec.get('columns'), spec.get('rows')
        if not isinstance(name, str) or not 1 <= len(name) <= 31 or re.search(r'[\\/*?:\[\]]', name) or name.lower() in names:
            raise ValueError('sheet names must be unique valid Excel names')
        if not isinstance(columns, list) or not 1 <= len(columns) <= 256 or not all(isinstance(c, str) and c for c in columns):
            raise ValueError('columns must be 1–256 nonempty header strings')
        if not isinstance(rows, list):
            raise ValueError('rows must be arrays matching columns')
        rows_total += len(rows)
        if rows_total > 100_000:
            raise ValueError('maximum 100000 rows per workbook')
        names.add(name.lower())
        sheet = book.create_sheet(name)
        for values in [columns, *rows]:
            if not isinstance(values, list) or len(values) != len(columns):
                raise ValueError('every row must match the column count')
            if any(v is not None and (type(v) not in (str, int, float, bool) or isinstance(v, str) and len(v) > 32767) for v in values):
                raise ValueError('cells must be scalar values of at most 32767 characters')
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # prevent formula injection from source data
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='245B82')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(14, max(len(str(c.value or '')) for c in column[:100]) + 2))
    fixed = datetime(2000, 1, 1)
    book.properties.created = book.properties.modified = fixed
    buffer = BytesIO()
    book.save(buffer)
    # Normalize ZIP timestamps and the library's current-time modified property.
    stable = BytesIO()
    with ZipFile(buffer) as original, ZipFile(stable, 'w', ZIP_DEFLATED) as target:
        for name in sorted(original.namelist()):
            value = original.read(name)
            if name == 'docProps/core.xml':
                value = re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)', rb'\g<1>2000-01-01T00:00:00Z\g<2>', value)
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, value)
    payload = stable.getvalue()
    checked = load_workbook(BytesIO(payload), read_only=True)
    try:
        if len(checked.sheetnames) != len(sheets):
            raise ValueError('reopen validation failed')
    finally:
        checked.close()
    dst.parent.mkdir(parents=True, exist_ok=True)
    import tempfile
    import os
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=dst.parent, suffix='.tmp', delete=False) as f:
            temporary = Path(f.name)
            f.write(payload)
            f.flush()
            os.fsync(f.fileno())
        temporary.replace(dst)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return {'path': output, 'sha256': sha256(payload).hexdigest(),
            'input_sha256': sha256(raw).hexdigest(), 'generator': VERSION,
            'rows': rows_total, 'sheets': len(sheets), 'validation_scope': 'schema and workbook structure'}


class GenerateWorkbookTool(Tool):
    name = 'generate_workbook'
    ends_turn_on_success = False
    description = ('Generate a validated Excel workbook from an intermediate JSON file. '
                   'Schema: {"sheets":[{"name":"Sheet1","columns":["Name","Qty"],"rows":[["Item",1]]}]}. '
                   'Literal cells only; no formulas. No shell required. Preserve source references in columns.')
    parameters = {'type': 'object', 'properties': {'source': {'type': 'string'}, 'output': {'type': 'string'}},
                  'required': ['source', 'output']}

    def __init__(self, workspace):
        self.workspace = workspace

    async def execute(self, source, output, **kwargs):
        from claw.jobs.provider import current_execution
        ctx = current_execution.get()
        # Trusted application generator in a cancellable child process, never
        # model-authored Python or a weaker shell sandbox.
        process = await asyncio.create_subprocess_exec(sys.executable, '-m', 'claw.jobs.workbook',
            str(self.workspace.resolve()), source, output,
            stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
        try:
            stdout, stderr = await process.communicate()
        except asyncio.CancelledError:
            if process.returncode is None:
                process.kill()
            await process.communicate()
            raise
        if process.returncode:
            from claw.jobs.tool_results import ToolText
            from claw.jobs.contracts import ToolResult
            return ToolText(ToolResult(status='error', error_code='invalid_workbook_input',
                text='Error: workbook generation failed: ' + stderr.decode('utf-8', 'replace')[-1000:]))
        result = json.loads(stdout)
        if ctx:
            await ctx.checkpoint({**ctx.state, 'generated': {**ctx.state.get('generated', {}), output: result}})
        return json.dumps(result, ensure_ascii=False)


if __name__ == '__main__':
    try:
        print(json.dumps(generate(*sys.argv[1:])))
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(1)
