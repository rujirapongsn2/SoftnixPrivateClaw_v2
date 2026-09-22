"""Safe ingestion and structured querying for Queryable Knowledge datasets."""

from __future__ import annotations

import asyncio
import csv
import datetime as dt
import json
import re
import shutil
import unicodedata
from itertools import zip_longest
from pathlib import Path
from typing import Any

import duckdb


MAX_SHEETS = 32
MAX_COLUMNS = 256
MAX_ROWS = 2_000_000
MAX_RESULT_ROWS = 200
_UNDERSCORES = re.compile(r"_+")


class DatasetError(ValueError):
    pass


def _identifier(value: Any, fallback: str, used: set[str]) -> str:
    raw = unicodedata.normalize("NFC", str(value or "").strip()).lower()
    base = "".join(
        character
        if character == "_" or unicodedata.category(character)[0] in {"L", "M", "N"}
        else "_"
        for character in raw
    )
    base = _UNDERSCORES.sub("_", base).strip("_") or fallback
    if base[0].isdigit():
        base = f"c_{base}"
    name = base
    index = 2
    while name in used:
        name = f"{base}_{index}"
        index += 1
    used.add(name)
    return name


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


def _is_number(value: str) -> bool:
    try:
        float(value.replace(",", ""))
        return True
    except ValueError:
        return False


def _is_datetime(value: str) -> bool:
    try:
        dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
        return True
    except ValueError:
        return False


def _column_profile(values: dict[str, int]) -> dict[str, Any]:
    non_empty = values["non_empty"]
    if non_empty and values["leading_zero"] == 0 and values["numbers"] == non_empty:
        inferred = "number"
    elif non_empty and values["datetimes"] == non_empty:
        inferred = "datetime"
    else:
        inferred = "text"
    return {
        "inferred_type": inferred,
        "non_empty_values": non_empty,
        "numeric_values": values["numbers"],
        "datetime_values": values["datetimes"],
    }


def _detect_text_encoding(path: Path) -> str:
    with path.open("rb") as stream:
        sample = stream.read(1_048_576)
    try:
        sample.decode("utf-8-sig")
        return "utf-8-sig"
    except UnicodeDecodeError:
        pass
    try:
        thai = sample.decode("cp874")
    except UnicodeDecodeError:
        thai = ""
    thai_chars = sum("\u0e00" <= char <= "\u0e7f" for char in thai)
    if thai_chars >= max(2, len(thai) // 100):
        return "cp874"
    return "cp1252"


def _quote(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _create_table(connection, table: str, headers: list[str], rows) -> tuple[int, list[dict]]:
    if not headers:
        raise DatasetError("the dataset has no header row")
    if len(headers) > MAX_COLUMNS:
        raise DatasetError(f"the dataset exceeds the {MAX_COLUMNS}-column limit")
    connection.execute(
        f"CREATE TABLE {_quote(table)} ({', '.join(f'{_quote(h)} VARCHAR' for h in headers)})"
    )
    placeholders = ", ".join("?" for _ in headers)
    insert = f"INSERT INTO {_quote(table)} VALUES ({placeholders})"
    profiles = [
        {"non_empty": 0, "numbers": 0, "datetimes": 0, "leading_zero": 0}
        for _ in headers
    ]
    batch: list[tuple[str | None, ...]] = []
    count = 0
    for raw in rows:
        values = tuple(_text(value) for value in list(raw)[: len(headers)])
        values += (None,) * (len(headers) - len(values))
        if not any(value not in (None, "") for value in values):
            continue
        count += 1
        if count > MAX_ROWS:
            raise DatasetError(f"the dataset exceeds the {MAX_ROWS:,}-row limit")
        for index, value in enumerate(values):
            if value in (None, ""):
                continue
            normalized = value.strip()
            profiles[index]["non_empty"] += 1
            profiles[index]["numbers"] += int(_is_number(normalized))
            profiles[index]["datetimes"] += int(_is_datetime(normalized))
            profiles[index]["leading_zero"] += int(
                len(normalized) > 1 and normalized[0] == "0" and normalized[1].isdigit()
            )
        batch.append(values)
        if len(batch) >= 1000:
            connection.executemany(insert, batch)
            batch.clear()
    if batch:
        connection.executemany(insert, batch)
    columns = [{"name": header, **_column_profile(profiles[index])} for index, header in enumerate(headers)]
    return count, columns


def _xlsx_rows(cached_sheet, formula_sheet, missing_formulas: list[int]):
    cached_rows = cached_sheet.iter_rows(values_only=True)
    formula_rows = formula_sheet.iter_rows(values_only=True)
    for cached, formulas in zip_longest(cached_rows, formula_rows, fillvalue=()):
        width = max(len(cached), len(formulas))
        values: list[Any] = []
        for index in range(width):
            cached_value = cached[index] if index < len(cached) else None
            formula_value = formulas[index] if index < len(formulas) else None
            if (
                isinstance(formula_value, str)
                and formula_value.startswith("=")
                and cached_value is None
            ):
                missing_formulas[0] += 1
            values.append(cached_value)
        yield tuple(values)


def ingest_dataset(source: Path, filename: str, destination: Path) -> tuple[dict[str, Any], int]:
    """Build a DuckDB index atomically and retain the immutable source file."""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise DatasetError("Queryable Knowledge supports CSV and XLSX files only")
    destination.mkdir(parents=True, exist_ok=True)
    source_copy = destination / f"source{suffix}"
    database = destination / "dataset.duckdb"
    temporary = destination / "dataset.tmp.duckdb"
    shutil.copy2(source, source_copy)
    temporary.unlink(missing_ok=True)
    tables: list[dict[str, Any]] = []
    total = 0
    connection = duckdb.connect(str(temporary))
    try:
        if suffix == ".csv":
            encoding = _detect_text_encoding(source)
            with source.open("r", encoding=encoding, errors="replace", newline="") as stream:
                sample = stream.read(65_536)
                stream.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.reader(stream, dialect)
                raw_headers = next(reader, None)
                if raw_headers is None:
                    raise DatasetError("the CSV file is empty")
                used: set[str] = set()
                headers = [
                    _identifier(value, f"column_{i + 1}", used)
                    for i, value in enumerate(raw_headers)
                ]
                count, columns = _create_table(connection, "data", headers, reader)
            tables.append({"name": "data", "source_name": filename, "rows": count, "columns": columns})
            total = count
        else:
            from openpyxl import load_workbook
            from claw.api.file_preview import _assert_workbook_within_budget

            _assert_workbook_within_budget(source)
            workbook = load_workbook(source, read_only=True, data_only=True)
            formula_workbook = None
            try:
                formula_workbook = load_workbook(source, read_only=True, data_only=False)
                if len(workbook.worksheets) > MAX_SHEETS:
                    raise DatasetError(f"the workbook exceeds the {MAX_SHEETS}-sheet limit")
                used_tables: set[str] = set()
                for sheet_index, sheet in enumerate(workbook.worksheets):
                    missing_formulas = [0]
                    formula_sheet = formula_workbook.worksheets[sheet_index]
                    iterator = _xlsx_rows(sheet, formula_sheet, missing_formulas)
                    raw_headers = next(iterator, None)
                    if raw_headers is None:
                        continue
                    used_columns: set[str] = set()
                    headers = [
                        _identifier(value, f"column_{i + 1}", used_columns)
                        for i, value in enumerate(raw_headers)
                    ]
                    table = _identifier(sheet.title, f"sheet_{sheet_index + 1}", used_tables)
                    count, columns = _create_table(connection, table, headers, iterator)
                    if missing_formulas[0]:
                        raise DatasetError(
                            f"sheet '{sheet.title}' contains {missing_formulas[0]} formula cell(s) "
                            "without cached values; open and save the workbook in Excel, then upload it again"
                        )
                    tables.append(
                        {"name": table, "source_name": sheet.title, "rows": count, "columns": columns}
                    )
                    total += count
            finally:
                workbook.close()
                if formula_workbook is not None:
                    formula_workbook.close()
        if not tables:
            raise DatasetError("the file contains no readable tables")
        connection.execute("CHECKPOINT")
    except Exception:
        connection.close()
        connection = None
        temporary.unlink(missing_ok=True)
        source_copy.unlink(missing_ok=True)
        raise
    finally:
        if connection is not None:
            connection.close()
    temporary.replace(database)
    schema = {"version": 1, "filename": filename, "tables": tables}
    (destination / "schema.json").write_text(json.dumps(schema, ensure_ascii=False), "utf-8")
    return schema, total


def run_query(database: Path, schema: dict[str, Any], request: dict[str, Any]) -> dict[str, Any]:
    tables = {table["name"]: table for table in schema.get("tables", [])}
    table_name = str(request.get("table") or "")
    if table_name not in tables:
        raise DatasetError(f"unknown table: {table_name}")
    table = tables[table_name]
    columns = {column["name"]: column for column in table.get("columns", [])}

    def column_expr(name: str, *, coercion: str = "text") -> str:
        if name not in columns:
            raise DatasetError(f"unknown column: {name}")
        quoted = _quote(name)
        if coercion == "number":
            return f"try_cast(replace({quoted}, ',', '') AS DECIMAL(38,6))"
        if coercion == "datetime":
            return f"try_cast({quoted} AS TIMESTAMP)"
        return quoted

    def kind(name: str) -> str:
        if name not in columns:
            raise DatasetError(f"unknown column: {name}")
        return str(columns[name].get("inferred_type") or "text")

    group_by = [str(value) for value in request.get("group_by") or []]
    selections = [column_expr(name) for name in group_by]
    headings = list(group_by)
    aggregations = request.get("aggregations") or []
    if not aggregations:
        selected = [str(value) for value in request.get("columns") or []] or list(columns)[:12]
        selections.extend(column_expr(name) for name in selected if name not in group_by)
        headings.extend(name for name in selected if name not in group_by)
    for index, item in enumerate(aggregations):
        function = str(item.get("function") or "").lower()
        name = str(item.get("column") or "")
        if function not in {"count", "count_distinct", "sum", "avg", "min", "max"}:
            raise DatasetError(f"unsupported aggregate: {function}")
        alias = str(item.get("alias") or f"{function}_{name or 'rows'}")
        if function == "count" and name in {"", "*"}:
            expression = "count(*)"
        elif function == "count_distinct":
            expression = f"count(DISTINCT {column_expr(name)})"
        elif function == "count":
            expression = f"count({column_expr(name)})"
        elif function in {"sum", "avg"}:
            if kind(name) != "number":
                profile = columns[name]
                raise DatasetError(
                    f"column '{name}' is not consistently numeric "
                    f"({profile.get('numeric_values', 0)} of {profile.get('non_empty_values', 0)} values); "
                    f"cannot calculate {function} safely"
                )
            expression = f"{function}({column_expr(name, coercion='number')})"
        else:
            profile = columns[name]
            non_empty = int(profile.get("non_empty_values", 0))
            numeric = int(profile.get("numeric_values", 0))
            datetimes = int(profile.get("datetime_values", 0))
            if kind(name) == "text" and (0 < numeric < non_empty or 0 < datetimes < non_empty):
                raise DatasetError(
                    f"column '{name}' contains mixed value types; cannot calculate {function} safely"
                )
            expression = f"{function}({column_expr(name, coercion=kind(name))})"
        selections.append(f"{expression} AS {_quote(alias)}")
        headings.append(alias)
    if not selections:
        raise DatasetError("select at least one column or aggregation")

    where: list[str] = []
    params: list[Any] = []
    for item in request.get("filters") or []:
        name = str(item.get("column") or "")
        operator = str(item.get("operator") or "eq")
        value = item.get("value")
        column_kind = kind(name)
        expression = column_expr(name, coercion=column_kind)
        if operator in {"eq", "ne", "gt", "gte", "lt", "lte"}:
            token = {"eq": "=", "ne": "!=", "gt": ">", "gte": ">=", "lt": "<", "lte": "<="}[operator]
            if operator in {"gt", "gte", "lt", "lte"} and column_kind == "text":
                raise DatasetError(
                    f"range comparison is unavailable for text column '{name}'; "
                    "clean the column or use eq, ne, contains, or starts_with"
                )
            if column_kind == "number":
                where.append(
                    f"{expression} {token} "
                    "try_cast(replace(cast(? AS VARCHAR), ',', '') AS DECIMAL(38,6))"
                )
            elif column_kind == "datetime":
                where.append(f"{expression} {token} try_cast(? AS TIMESTAMP)")
            else:
                where.append(f"{expression} {token} ?")
            params.append(value)
        elif operator in {"contains", "starts_with"}:
            where.append(f"lower({column_expr(name)}) LIKE lower(?)")
            params.append(f"%{value}%" if operator == "contains" else f"{value}%")
        else:
            raise DatasetError(f"unsupported filter operator: {operator}")

    sql = f"SELECT {', '.join(selections)} FROM {_quote(table_name)}"
    if where:
        sql += " WHERE " + " AND ".join(where)
    if group_by:
        sql += " GROUP BY " + ", ".join(column_expr(name) for name in group_by)
    limit = max(1, min(int(request.get("limit") or 50), MAX_RESULT_ROWS))
    sql += f" LIMIT {limit}"
    connection = duckdb.connect(str(database), read_only=True)
    try:
        connection.execute("SET threads = 1")
        connection.execute("SET memory_limit = '512MB'")
        cursor = connection.execute(sql, params)
        rows = cursor.fetchall()
    finally:
        connection.close()
    return {
        "table": table_name,
        "columns": headings,
        "rows": [[str(value) if value is not None else None for value in row] for row in rows],
        "returned_rows": len(rows),
        "source_rows": table.get("rows", 0),
    }


async def execute_dataset_query(
    store,
    root: Path,
    user_id: str,
    *,
    knowledge_base: str,
    action: str = "inspect",
    dataset: str = "",
    table: str = "",
    columns: list[str] | None = None,
    filters: list[dict[str, Any]] | None = None,
    aggregations: list[dict[str, Any]] | None = None,
    group_by: list[str] | None = None,
    limit: int = 50,
) -> str:
    bases = [base for base in await store.list_accessible(user_id) if base.get("kind") == "queryable"]
    needle = knowledge_base.strip().lower()
    matches = [base for base in bases if needle == base["name"].lower()]
    if not matches:
        matches = [base for base in bases if needle in base["name"].lower()]
    if not matches:
        return json.dumps({"error": "Queryable Knowledge not found"}, ensure_ascii=False)
    if len(matches) > 1:
        return json.dumps(
            {"error": "Knowledge name is ambiguous", "matches": [base["name"] for base in matches]},
            ensure_ascii=False,
        )
    base = matches[0]
    docs = [doc for doc in await store.list_docs(base["id"]) if doc.status == "ready" and doc.dataset_path]
    if dataset:
        docs = [doc for doc in docs if dataset.lower() in doc.filename.lower()]
    if not docs:
        return json.dumps({"error": "No ready dataset is available"}, ensure_ascii=False)
    if action not in {"inspect", "query"}:
        return json.dumps({"error": f"Unsupported action: {action}"}, ensure_ascii=False)
    if action == "inspect":
        return json.dumps(
            {
                "knowledge_base": base["name"],
                "datasets": [
                    {
                        "filename": doc.filename,
                        "rows": doc.dataset_rows,
                        "tables": (doc.dataset_schema or {}).get("tables", []),
                    }
                    for doc in docs
                ],
            },
            ensure_ascii=False,
        )
    if len(docs) > 1 and not dataset:
        return json.dumps(
            {"error": "Choose a dataset", "datasets": [doc.filename for doc in docs]},
            ensure_ascii=False,
        )
    doc = docs[0]
    root = root.resolve()
    database = (root / doc.dataset_path).resolve()
    try:
        database.relative_to(root)
    except ValueError:
        return json.dumps({"error": "Invalid dataset path"})
    if not database.is_file():
        return json.dumps({"error": "Dataset index is missing"})
    try:
        result = await asyncio.to_thread(
            run_query,
            database,
            doc.dataset_schema or {},
            {
                "table": table,
                "columns": columns or [],
                "filters": filters or [],
                "aggregations": aggregations or [],
                "group_by": group_by or [],
                "limit": limit,
            },
        )
    except DatasetError as exc:
        return json.dumps({"error": str(exc)}, ensure_ascii=False)
    result.update(
        {
            "knowledge_base": base["name"],
            "dataset": doc.filename,
            "dataset_id": doc.id,
        }
    )
    return json.dumps(result, ensure_ascii=False)
