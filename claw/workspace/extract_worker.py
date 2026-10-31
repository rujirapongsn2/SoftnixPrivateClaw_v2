"""Disposable worker that extracts the searchable text of one workspace file.

Workspace files are untrusted: the agent's sandbox or the user can write a
crafted PDF or a zip bomb disguised as .docx/.xlsx. Parsing happens here, in a
child process with CPU, memory and file limits, so a hostile file can at worst
kill this process; :mod:`claw.workspace.content_index` also kills it on a
timeout. The text goes to stdout, already capped.

Usage: python -m claw.workspace.extract_worker <path>
"""

import mimetypes
import resource
import sys
from pathlib import Path

MAX_TEXT_CHARS = 100_000
_MEMORY_BYTES = 1024 * 1024 * 1024


def _limits() -> None:
    limits = (
        (resource.RLIMIT_CORE, 0, 0),
        (resource.RLIMIT_CPU, 20, 25),
        (resource.RLIMIT_FSIZE, 0, 0),
        (resource.RLIMIT_AS, _MEMORY_BYTES, _MEMORY_BYTES),
    )
    for kind, soft, hard in limits:
        try:
            resource.setrlimit(kind, (soft, hard))
        except (ValueError, OSError):
            pass  # not every platform enforces every limit (e.g. RLIMIT_AS on macOS)


def _xlsx_text(source: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=True)
    parts: list[str] = []
    size = 0
    try:
        for sheet in workbook.worksheets:
            parts.append(sheet.title)
            for row in sheet.iter_rows(values_only=True):
                line = " ".join(str(v) for v in row if v is not None)
                if line:
                    parts.append(line)
                    size += len(line)
                if size >= MAX_TEXT_CHARS:
                    return "\n".join(parts)
    finally:
        workbook.close()
    return "\n".join(parts)


def extract(source: Path) -> str:
    """Searchable text of one file, capped; "" when it cannot be read."""
    from claw.knowledge.parse import extract_text

    try:
        if source.suffix.lower() == ".xlsx":
            text = _xlsx_text(source)
        else:
            mime = mimetypes.guess_type(source.name)[0] or ""
            text, _error = extract_text(source.read_bytes(), source.name, mime)
    except Exception:  # noqa: BLE001 - an unreadable file is indexed as empty
        return ""
    return text[:MAX_TEXT_CHARS]


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(2)
    _limits()
    sys.stdout.write(extract(Path(sys.argv[1])))


if __name__ == "__main__":
    main()
