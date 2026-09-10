"""Opt-in real-model document chain in a fresh database with synthetic input only."""
import asyncio
import json
import sys
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from dotenv import load_dotenv
from loguru import logger
from docx import Document
from claw.config import Settings as ClawSettings
from sbot.config import Settings, LLMSettings, SandboxSettings
from sbot.core.missions import MissionService
from sbot.db.engine import create_engine_and_factory, init_db
from sbot.db.stores import UserStore, BotStore, MissionStore
from sbot.providers.litellm_provider import LiteLLMProvider
from sbot.providers.base import ChatResult
from sbot.sandbox.ephemeral import EphemeralSandbox

async def main():
    logger.remove()
    if len(sys.argv) > 1:
        load_dotenv(sys.argv[1])
    config = ClawSettings()
    # Resolve the admin default exactly as specialist dispatch does; env-only
    # credentials may refer to an obsolete provider when Control Plane is used.
    from claw.db.stores import LLMConfigStore
    from claw.security.crypto import SecretBox
    config_engine, config_factory = create_engine_and_factory(config.database_url)
    registry = LLMConfigStore(config_factory, SecretBox(config.secret_key))
    default_id = await registry.default_model_for(None)
    resolved = await registry.resolve(default_id) if default_id else None
    await config_engine.dispose()
    if resolved:
        config.llm.model = resolved['model_id']
        config.llm.api_key = resolved['api_key']
        config.llm.api_base = resolved['api_base'] or ''
    root = Path(tempfile.mkdtemp(prefix='team-document-e2e-'))
    settings = Settings(_env_file=None, database_url=f'sqlite+aiosqlite:///{root}/test.db',
        workspaces_root=root/'workspaces', llm=LLMSettings(model=config.llm.model),
        sandbox=SandboxSettings(enabled=True, network='none', timeout_seconds=60))
    class Metered(LiteLLMProvider):
        calls = 0
        tokens = 0
        async def stream_chat(self, *args, **kwargs):
            self.calls += 1
            if self.calls > 65:
                raise RuntimeError('Test call budget reached')
            async for event in super().stream_chat(*args, **kwargs):
                if isinstance(event, ChatResult):
                    self.tokens += sum(event.usage.get(k, 0) for k in ('prompt_tokens','completion_tokens'))
                    print(json.dumps({'call': self.calls, 'finish': event.finish_reason, 'tokens': self.tokens}), flush=True)
                yield event
    provider = Metered(api_key=config.llm.api_key, api_base=config.llm.api_base, default_model=config.llm.model)
    engine, factory = create_engine_and_factory(settings.database_url)
    await init_db(engine)
    users, bots, missions = UserStore(factory), BotStore(factory), MissionStore(factory)
    service = MissionService(missions, bots, provider, EphemeralSandbox(settings.sandbox), settings)
    report = {'root': str(root), 'model': config.llm.model}
    try:
        user = await users.get_or_create_by_email('doc-chain@example.invalid')
        leader = await bots.get_or_create_cos(user.id)
        analyst = await bots.create(owner_id=user.id, name='Analyst', role_title='Document reviewer', charter='Review facts from files. Save concise findings with evidence.')
        writer = await bots.create(owner_id=user.id, name='Writer', role_title='Document editor', charter='Edit Word documents precisely, preserving unrelated content and tables.')
        workspace = settings.workspaces_root/user.id
        workspace.mkdir(parents=True)
        d = Document()
        for i in range(1, 61):
            d.add_heading(f'ข้อกำหนด {i}', 1)
            d.add_paragraph(f'ระบบส่วนที่ {i} ต้องบันทึกข้อมูลพร้อมตรวจสอบย้อนกลับและกำหนดสิทธิ์ผู้ใช้งานตามบทบาท ห้ามเปลี่ยนข้อมูลส่วนนี้โดยไม่มีเหตุผล ' * 4)
        t = d.add_table(rows=2, cols=2)
        t.cell(0,0).text='Metric';t.cell(0,1).text='Requirement'
        t.cell(1,0).text='EPS';t.cell(1,1).text='9000'
        d.save(workspace/'draft.docx')
        (workspace/'component.txt').write_text('Verified component capacity: EPS 5000. Correct the draft capacity to 5000. All other paragraphs and tables must be preserved.')
        job = await service.submit_work(user.id, None, 'document-chain', 'Review the long Thai Word specification against component.txt, then correct the Word document and provide a change log.', [
            {'id':'review','bot_id':analyst.id,'instruction':'Inspect draft.docx including its table against component.txt. Save review.md identifying the discrepancy and exact correction with source reference. Do not rewrite the document.', 'input_files':['draft.docx','component.txt'],'required_files':['review.md']},
            {'id':'rewrite','bot_id':writer.id,'depends_on':['review'],'instruction':'Read upstream review.md and original draft.docx from the snapshot manifest. Use python-docx to apply the correction in the existing document and save revised.docx. Preserve all 120 paragraphs and the table. Reopen to verify the EPS cell is 5000, then write changes.md. Finish with both files.', 'input_files':['draft.docx'],'required_files':['revised.docx','changes.md']},
        ], leader.id)
        async with asyncio.timeout(480):
            report['status'] = await service._running[job.id]
        outputs = list(workspace.rglob('revised.docx'))
        final = Document(outputs[0]) if outputs else None
        report['checks'] = {'completed':report['status']=='completed', 'paragraphs_preserved':final is not None and len(final.paragraphs)==120, 'table_corrected':final is not None and len(final.tables)==1 and final.tables[0].cell(1,1).text=='5000', 'original_unchanged':Document(workspace/'draft.docx').tables[0].cell(1,1).text=='9000'}
    except Exception as exc:
        report['error_type'] = type(exc).__name__
    finally:
        report.update(calls=provider.calls,tokens=provider.tokens)
        await service.stop()
        await engine.dispose()
        (root/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
        print(json.dumps(report,ensure_ascii=False),flush=True)

asyncio.run(main())
