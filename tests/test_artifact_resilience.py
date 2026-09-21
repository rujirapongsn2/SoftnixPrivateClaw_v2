"""Regression cases from the failed TOR-to-BOM workflow."""
import json
from unittest.mock import AsyncMock

from docx import Document

from claw.sandbox.ephemeral import SandboxResult
from claw.tools.documents import ReadDocxTool
from claw.providers.base import ChatResult, ToolCall
from tests.conftest import FakeProvider
from tests.test_artifact_jobs import make_runtime
from claw.db.models import AppSetting


def test_shell_failures_are_errors_not_successes():
    assert SandboxResult(1, '', 'bad script').render().startswith('Error:')
    assert SandboxResult(125, '', 'missing image', infrastructure_error=True).render().startswith(
        'Error: [sandbox_unavailable]')
    assert not SandboxResult(0, 'ok', '').render().startswith('Error:')


async def test_docx_pages_preserve_tables_and_tail_without_shell(tmp_path):
    doc = Document()
    doc.add_paragraph('START')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Hardware'
    table.cell(0, 1).text = '42 servers'
    doc.add_paragraph('long TOR clause ' * 4000)
    doc.add_paragraph('END-SERVICES')
    doc.save(tmp_path / 'tor.docx')
    reader = ReadDocxTool(tmp_path)
    offset, chunks = 0, []
    while offset is not None:
        page = json.loads(await reader.execute('tor.docx', offset=offset))
        chunks.append(page['text'])
        offset = page['next_offset']
    text = ''.join(chunks)
    assert len(chunks) > 5
    assert text.index('START') < text.index('Hardware') < text.index('42 servers')
    assert text.endswith('END-SERVICES')
    assert len(reader._parsed) == 1


async def test_infrastructure_failure_stops_model_loop_and_retains_checkpoint(
    stores, db_factory, tmp_path
):
    call = ToolCall(id='exec-1', name='exec', arguments={'command': 'python make.py'})
    provider = FakeProvider([[ChatResult(content=None, tool_calls=[call])], [ChatResult(content='wrong')]])
    runtime, jobs = make_runtime(stores, db_factory, provider, tmp_path)
    runtime.sandbox.run = AsyncMock(return_value=SandboxResult(
        125, '', 'No such image', infrastructure_error=True))
    user = await stores['users'].get_or_create_by_email('blocked@example.com')
    session = await stores['sessions'].create(user.id)
    result = await runtime.handle_message(user.id, session.id, 'สร้างไฟล์ Excel', locale='th')
    assert 'พักงาน' in result
    assert len(provider.calls) == 1
    saved = (await jobs.list_for_session(user.id, session.id, active_only=False))[0]
    assert saved['status'] == 'blocked'
    assert saved['checkpoint_messages']
    assert saved['content'] == 'สร้างไฟล์ Excel'
    assert not saved['tool_results']  # failed exec is never cached as successful
    assert runtime.sandbox.run.await_count == 1
    await runtime.drain()


async def test_continue_uses_same_checkpoint_after_sandbox_restored(stores, db_factory, tmp_path):
    call = ToolCall(id='exec-1', name='exec', arguments={'command': 'python make.py'})
    provider = FakeProvider([[ChatResult(content=None, tool_calls=[call])], [ChatResult(content='Done')]])
    runtime, jobs = make_runtime(stores, db_factory, provider, tmp_path)
    runtime.sandbox.run = AsyncMock(return_value=SandboxResult(125, '', 'missing', infrastructure_error=True))
    user = await stores['users'].get_or_create_by_email('resume-blocked@example.com')
    session = await stores['sessions'].create(user.id)
    await runtime.handle_message(user.id, session.id, 'Create Excel')
    runtime.sandbox.run = AsyncMock(return_value=SandboxResult(0, 'ready', ''))
    assert await runtime.handle_message(user.id, session.id, 'ทำต่อให้จบ') == 'Done'
    saved = await jobs.list_for_session(user.id, session.id, active_only=False)
    assert len(saved) == 1
    assert saved[0]['status'] == 'completed'
    assert 'sandbox_unavailable' in str(provider.calls[-1])
    await runtime.drain()


async def test_admin_policy_is_snapshotted_for_normal_chat_job(stores, db_factory, tmp_path):
    async with db_factory() as db:
        db.add(AppSetting(key='team_resource_policy', value={'policy': {
            'max_job_tokens': 7777, 'max_job_seconds': 1234,
            'automatic_resources': True, 'max_resource_adjustments': 2,
            'max_step_recoveries': 1,
        }}))
        await db.commit()
    runtime, jobs = make_runtime(stores, db_factory, FakeProvider([[ChatResult(content='Done')]]), tmp_path)
    user = await stores['users'].get_or_create_by_email('policy-artifact@example.com')
    session = await stores['sessions'].create(user.id)
    await runtime.handle_message(user.id, session.id, 'Create Excel')
    job = (await jobs.list_for_session(user.id, session.id, active_only=False))[0]
    assert job['budget'] == {'seconds': 1234, 'tokens': 7777, 'extensions': 2, 'recoveries': 1}
    await runtime.drain()


async def test_runtime_recovers_dependency_without_repeating_model_retries(stores, db_factory, tmp_path):
    first = ToolCall(id='first', name='exec', arguments={'command': 'python make.py'})
    retry = ToolCall(id='second', name='exec', arguments={'command': 'python make.py'})
    provider = FakeProvider([
        [ChatResult(content=None, tool_calls=[first])],
        [ChatResult(content=None, tool_calls=[retry])],
        [ChatResult(content='Done')],
    ])
    runtime, jobs = make_runtime(stores, db_factory, provider, tmp_path,
                                 artifact_job_max_seconds=60, max_turn_seconds=1)
    runtime.sandbox.run = AsyncMock(side_effect=[
        SandboxResult(125, '', 'missing', infrastructure_error=True),
        SandboxResult(0, '', ''),  # runtime health probe
        SandboxResult(0, 'generated', ''),
    ])
    user = await stores['users'].get_or_create_by_email('auto-recover@example.com')
    session = await stores['sessions'].create(user.id)
    assert await runtime.handle_message(user.id, session.id, 'Create Excel') == 'Done'
    assert [c.args[0] for c in runtime.sandbox.run.await_args_list] == [
        'python make.py', 'true', 'python make.py']
    job = (await jobs.list_for_session(user.id, session.id, active_only=False))[0]
    assert job['status'] == 'completed'
    assert job['elapsed_seconds'] >= 5
    await runtime.drain()
