"""Real normal runtime + durable worker, scripted provider, real filesystem/DB."""
from functools import partial

from claw.jobs.store import JobStore
from claw.jobs.worker import Worker, Executor
from claw.jobs.provider import AccountedProvider
from claw.jobs.runtime_bridge import normal_execute, file_evidence
from claw.security.crypto import SecretBox
from claw.providers.base import ChatResult, ToolCall
from tests.conftest import FakeProvider
from tests.test_artifact_jobs import make_runtime


async def prepare(stores, db_factory, tmp_path, turns):
    provider = FakeProvider(turns)
    runtime, _ = make_runtime(stores, db_factory, AccountedProvider(provider), tmp_path,
                              max_turn_seconds=10, artifact_job_max_tokens=500000)
    runtime.settings.durable_jobs_privateclaw = True
    jobs = JobStore(db_factory, SecretBox('test'))
    await jobs.initialize()
    runtime.durable_jobs = jobs
    user = await stores['users'].get_or_create_by_email('worker-runtime@example.test')
    session = await stores['sessions'].create(user.id)

    async def allow(*args):
        return True

    async def validate(ctx, result):
        if result.status == 'yielded':
            from claw.jobs.runtime_bridge import validate_progress
            return validate_progress(runtime.get_agent(user.id).workspace, ctx.state, result.evidence)
        if ctx.lease.spec.acceptance.get('kind') == 'research':
            from claw.jobs.research import research_evidence
            return bool(result.delivery) and research_evidence(ctx.state, result.delivery['content']) == result.evidence
        refs = result.evidence.get('files', [])
        return bool(refs) and file_evidence(runtime.get_agent(user.id).workspace, [r['path'] for r in refs]) == result.evidence

    worker = Worker(jobs, {'privateclaw.turn': Executor(partial(normal_execute, runtime), validate)}, allow)
    return runtime, jobs, worker, user, session, provider


def test_yield_progress_requires_current_workspace_evidence(tmp_path):
    from claw.jobs.runtime_bridge import progress_evidence, validate_progress

    output = tmp_path / 'result.csv'
    output.write_text('name,value\na,1\n')
    state = {'runtime': {'written': ['result.csv']}}
    evidence = progress_evidence(tmp_path, state)
    assert evidence['files'][0]['path'] == 'result.csv'
    assert validate_progress(tmp_path, state, evidence)

    output.write_text('name,value\na,2\n')
    assert not validate_progress(tmp_path, state, evidence)
    assert not validate_progress(tmp_path, {'runtime': {'written': ['missing.csv']}},
                                 {'coverage': {}, 'research': {}, 'files': []})


async def test_real_runtime_admits_without_provider_and_delivers_once(stores, db_factory, tmp_path):
    values = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content=None, tool_calls=[ToolCall('write', 'write_file', {'path': 'result.csv', 'content': 'name,value\na,1\n'})])],
        [ChatResult(content='สร้างไฟล์เรียบร้อย')],
    ])
    runtime, jobs, worker, user, session, provider = values
    await runtime.handle_message(user.id, session.id, 'ช่วยสร้างไฟล์ CSV', permission_mode='auto')
    assert not provider.calls
    assert len(await jobs.list_jobs(user.id, session.id)) == 1
    assert await worker.run_once()
    job = (await jobs.list_jobs(user.id, session.id))[0]
    assert job['status'] == 'completed', job
    assert await jobs.publish_deliveries() == 1
    assert await jobs.publish_deliveries() == 0
    from sqlalchemy import select
    from claw.db.models import Message
    async with db_factory() as db:
        rows = list((await db.scalars(select(Message).where(Message.session_id == session.id).order_by(Message.seq))).all())
    assert sum(r.role == 'user' for r in rows) == 1
    assert sum(r.role == 'assistant' for r in rows) == 1
    assert rows[-1].meta['artifacts'] == ['result.csv']


async def test_missing_file_cannot_be_success(stores, db_factory, tmp_path):
    runtime, jobs, worker, user, session, _ = await prepare(stores, db_factory, tmp_path,
        [[ChatResult(content='สร้างไฟล์เรียบร้อย')]])
    await runtime.handle_message(user.id, session.id, 'ช่วยสร้างไฟล์ CSV')
    await worker.run_once()
    job = (await jobs.list_jobs(user.id, session.id))[0]
    assert job['status'] == 'paused'
    assert job['reason'] == 'missing_deliverable'
    assert await jobs.publish_deliveries() == 0


async def test_dependency_recovery_does_not_call_model_while_waiting(stores, db_factory, tmp_path):
    from unittest.mock import AsyncMock
    from claw.sandbox.ephemeral import SandboxResult
    runtime, jobs, worker, user, session, provider = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content=None, tool_calls=[ToolCall('exec', 'exec', {'command': 'echo test'})])],
        [ChatResult(content=None, tool_calls=[ToolCall('write', 'write_file', {'path': 'result.csv', 'content': 'a,b\n1,2'})])],
        [ChatResult(content='Done')],
    ])
    runtime.sandbox.run = AsyncMock(return_value=SandboxResult(125, '', 'not ready', infrastructure_error=True))
    await runtime.handle_message(user.id, session.id, 'Create a CSV file', permission_mode='auto')
    await worker.run_once()
    job = (await jobs.list_jobs(user.id, session.id))[0]
    assert job['status'] == 'waiting_dependency'
    assert len(provider.calls) == 1
    assert not await worker.run_once()
    assert len(provider.calls) == 1
    from claw.jobs.models import Step
    async with db_factory() as db:
        step = await db.get(Step, (job['job_id'], 'task'))
        step.next_at = 0
        await db.commit()
    worker.probes[('privateclaw', 'sandbox')] = AsyncMock(return_value=True)
    await worker.probe_once()
    await worker.run_once()
    assert (await jobs.snapshot(user.id, job['job_id']))['status'] == 'completed'
    assert runtime.sandbox.run.await_count == 1


async def test_candidate_recovery_only_validates_never_calls_model(stores, db_factory, tmp_path):
    from dataclasses import asdict
    from claw.jobs.contracts import StepOutcome
    runtime, jobs, worker, user, session, provider = await prepare(stores, db_factory, tmp_path, [])
    await runtime.handle_message(user.id, session.id, 'Create CSV file')
    lease = await jobs.claim('crashed', {'privateclaw.turn'})
    path = runtime.get_agent(user.id).workspace / 'result.csv'
    path.write_text('a,b\n1,2')
    evidence = file_evidence(path.parent, ['result.csv'])
    candidate = StepOutcome('completed', evidence=evidence,
        delivery={'content': 'Done', 'meta': {'artifacts': ['result.csv']}})
    await jobs.checkpoint(lease, {'candidate': asdict(candidate)})
    from claw.jobs.models import Step
    async with db_factory() as db:
        step = await db.get(Step, (lease.job_id, lease.step_id))
        step.lease_until = 0
        await db.commit()
    await worker.run_once()
    assert not provider.calls
    assert (await jobs.snapshot(user.id, lease.job_id))['status'] == 'completed'
    assert await jobs.publish_deliveries() == 1


async def test_docx_cache_survives_new_reader_and_tracks_complete_coverage(stores, db_factory, tmp_path):
    from unittest.mock import patch
    from docx import Document
    from claw.jobs.provider import current_execution
    from claw.jobs.worker import ExecutionContext
    from claw.jobs.runtime_bridge import coverage_complete
    from claw.tools.documents import ReadDocxTool
    runtime, jobs, _, user, session, _ = await prepare(stores, db_factory, tmp_path, [])
    root = runtime.get_agent(user.id).workspace
    doc = Document()
    doc.add_paragraph('x' * 5000)
    doc.save(root / 'source.docx')
    await runtime.handle_message(user.id, session.id, 'Create Excel file', media=['source.docx'])
    lease = await jobs.claim('test', {'privateclaw.turn'})
    ctx = ExecutionContext(jobs, lease)
    token = current_execution.set(ctx)
    try:
        original = ReadDocxTool._extract
        with patch.object(ReadDocxTool, '_extract', wraps=original) as extract:
            await ReadDocxTool(root).execute('source.docx')
            assert not coverage_complete(ctx.state, ['source.docx'])
            await ReadDocxTool(root).execute('source.docx', offset=4000)
            assert extract.call_count == 1
            assert coverage_complete(ctx.state, ['source.docx'])
    finally:
        current_execution.reset(token)


async def test_business_approval_survives_worker_boundary(stores, db_factory, tmp_path):
    runtime, jobs, worker, user, session, provider = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content=None, tool_calls=[ToolCall('exec1', 'exec', {'command': 'echo ok'})])],
        [ChatResult(content=None, tool_calls=[ToolCall('exec2', 'exec', {'command': 'echo ok'})])],
        [ChatResult(content=None, tool_calls=[ToolCall('write', 'write_file', {'path': 'result.csv', 'content': 'a,b'})])],
        [ChatResult(content='Done')],
    ])
    from unittest.mock import AsyncMock
    from claw.sandbox.ephemeral import SandboxResult
    runtime.sandbox.run = AsyncMock(return_value=SandboxResult(0, 'ok', ''))
    await runtime.handle_message(user.id, session.id, 'Create a CSV file', permission_mode='ask')
    await worker.run_once()
    job = (await jobs.list_jobs(user.id, session.id))[0]
    assert job['status'] == 'awaiting_input'
    request = job['steps'][0]['approval']
    assert not runtime.sandbox.run.called
    assert not await jobs.approve('other-owner', job['job_id'], 'task', request['key'], True)
    assert await jobs.approve(user.id, job['job_id'], 'task', request['key'], True)
    assert not await jobs.approve(user.id, job['job_id'], 'task', request['key'], True)
    await worker.run_once()
    assert runtime.sandbox.run.await_count == 1
    assert (await jobs.snapshot(user.id, job['job_id']))['status'] == 'completed'


def test_workbook_generator_deterministic_literal_cells_and_schema(tmp_path):
    import json
    import pytest
    from openpyxl import load_workbook
    from claw.jobs.workbook import generate
    source = tmp_path / 'data.json'
    source.write_text(json.dumps({'sheets': [{'name': 'Data', 'columns': ['Name', 'Qty'],
        'rows': [['=HYPERLINK("https://example.test")', 2], ['ไทย', 1]]}]}))
    one = generate(tmp_path, 'data.json', 'a.xlsx')
    two = generate(tmp_path, 'data.json', 'b.xlsx')
    assert one['sha256'] == two['sha256']
    book = load_workbook(tmp_path / 'a.xlsx')
    assert book['Data']['A2'].data_type == 's'
    book.close()
    source.write_text(json.dumps({'sheets': [{'name': 'Data', 'columns': ['A'], 'rows': [[1, 2]]}]}))
    with pytest.raises(ValueError, match='column count'):
        generate(tmp_path, 'data.json', 'invalid.xlsx')
    assert not (tmp_path / 'invalid.xlsx').exists()


async def test_xlsx_runtime_uses_trusted_generator_without_shell(stores, db_factory, tmp_path):
    import json
    from unittest.mock import AsyncMock
    runtime, jobs, worker, user, session, _ = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content=None, tool_calls=[ToolCall('write', 'write_file', {'path': 'data.json',
            'content': json.dumps({'sheets': [{'name': 'Data', 'columns': ['Name'], 'rows': [['Test']]}]})})])],
        [ChatResult(content=None, tool_calls=[ToolCall('generate', 'generate_workbook', {'source': 'data.json', 'output': 'result.xlsx'})])],
        [ChatResult(content='Done')],
    ])
    runtime.sandbox.run = AsyncMock(side_effect=AssertionError('must not need shell'))
    await runtime.handle_message(user.id, session.id, 'Create Excel file')
    await worker.run_once()
    job = (await jobs.list_jobs(user.id, session.id))[0]
    assert job['status'] == 'completed', job
    assert not runtime.sandbox.run.called


async def test_multisource_research_finishes_in_chat_without_artifact(stores, db_factory, tmp_path, monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    fetch = AsyncMock(return_value=SimpleNamespace(status_code=200, headers={'content-type': 'text/plain'}, text='Verified retrieval fixture'))
    monkeypatch.setattr('httpx.AsyncClient.get', fetch)
    runtime, jobs, worker, user, session, provider = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content=None, tool_calls=[ToolCall('s1', 'web_fetch', {'url': 'https://one.test/source'})])],
        [ChatResult(content=None, tool_calls=[ToolCall('s2', 'web_fetch', {'url': 'https://two.test/source'})])],
        [ChatResult(content='Summary with [one](https://one.test/source) and [two](https://two.test/source).')],
    ])
    await runtime.handle_message(user.id, session.id, 'Research and compare the two sources, then summarize findings.')
    assert not provider.calls  # admission, not foreground execution
    await worker.run_once()
    job = (await jobs.list_jobs(user.id, session.id))[0]
    assert job['status'] == 'completed', job
    assert fetch.await_count == 2
    assert await jobs.publish_deliveries() == 1
    assert await jobs.publish_deliveries() == 0
    messages = await stores['messages'].recent(session.id)
    result = next(m for m in messages if m['role'] == 'assistant')
    assert 'factual completeness not verified' in result['meta']['validation_scope']


async def test_research_cannot_complete_with_invented_citations(stores, db_factory, tmp_path):
    runtime, jobs, worker, user, session, _ = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content='Complete: [source](https://invented.test/page)')],
    ])
    await runtime.handle_message(user.id, session.id, 'Research several sources and summarize findings.')
    await worker.run_once()
    job = (await jobs.list_jobs(user.id, session.id))[0]
    assert job['status'] == 'paused' and job['reason'] == 'validation_failed'
    assert await jobs.publish_deliveries() == 0


async def test_research_source_cache_survives_new_worker_claim(stores, db_factory, tmp_path, monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    from claw.jobs.provider import current_execution
    from claw.jobs.worker import ExecutionContext
    from claw.jobs.contracts import StepOutcome
    from claw.tools.web import WebFetchTool
    runtime, jobs, _, user, session, _ = await prepare(stores, db_factory, tmp_path, [])
    fetch = AsyncMock(return_value=SimpleNamespace(status_code=200, headers={'content-type': 'text/plain'}, text='Retained source'))
    monkeypatch.setattr('httpx.AsyncClient.get', fetch)
    await runtime.handle_message(user.id, session.id, 'Research multiple sources and summarize findings.')
    first = await jobs.claim('first', {'privateclaw.turn'})
    ctx = ExecutionContext(jobs, first)
    token = current_execution.set(ctx)
    try:
        original = await WebFetchTool().execute('https://source.test/page')
    finally:
        current_execution.reset(token)
    await jobs.settle(first, StepOutcome('yielded', checkpoint=ctx.state))
    second = await jobs.claim('replacement', {'privateclaw.turn'})
    token = current_execution.set(ExecutionContext(jobs, second))
    try:
        assert await WebFetchTool().execute('https://source.test/page') == original
    finally:
        current_execution.reset(token)
    assert fetch.await_count == 1


def test_research_negated_file_instruction_and_source_urls_are_not_artifact_requests():
    from claw.core.runtime import _is_artifact_task
    from claw.jobs.research import is_multistep_research
    instruction = ('Research and summarize Python lists using https://docs.python.org/3/tutorial/datastructures.html '
                   'and https://docs.python.org/3/library/stdtypes.html. Do not create files.')
    assert is_multistep_research(instruction)
    assert not _is_artifact_task(instruction)
    assert not _is_artifact_task('ค้นคว้าและสรุปหลายแหล่ง ไม่ต้องสร้างไฟล์')
    assert _is_artifact_task('Research and summarize sources in an Excel file. Create the workbook.')


async def test_foreground_promotes_without_replaying_completed_tools(stores, db_factory, tmp_path):
    """A newly discovered file workflow keeps prior writes and billed calls."""
    from unittest.mock import AsyncMock
    from sqlalchemy import select
    from claw.jobs.models import ResourceCall, Step
    from claw.db.models import Message, UsageRecord
    plan = {'goal': 'Create a CSV file with the analysis', 'steps': [
        {'step': 'Prepare input', 'status': 'done'},
        {'step': 'Generate output', 'status': 'in_progress'},
        {'step': 'Check output', 'status': 'pending'}]}
    runtime, jobs, worker, user, session, provider = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content=None, tool_calls=[ToolCall('prepare', 'write_file', {'path': 'prepared.txt', 'content': 'ready'})],
                    usage={'prompt_tokens': 100, 'completion_tokens': 20})],
        [ChatResult(content=None, tool_calls=[ToolCall('plan', 'update_plan', plan)],
                    usage={'prompt_tokens': 150, 'completion_tokens': 30})],
        [ChatResult(content=None, tool_calls=[ToolCall('output', 'write_file', {'path': 'result.csv', 'content': 'name,value\na,1'})],
                    usage={'prompt_tokens': 80, 'completion_tokens': 10})],
        [ChatResult(content='Done', usage={'prompt_tokens': 50, 'completion_tokens': 5})],
    ])
    tool = runtime.get_agent(user.id).tools.get('write_file')
    original = tool.execute
    tool.execute = AsyncMock(side_effect=original)
    await runtime.handle_message(user.id, session.id, 'Please process this analysis', permission_mode='auto')
    assert len(provider.calls) == 2  # foreground yielded immediately after plan checkpoint
    job = (await jobs.list_jobs(user.id, session.id))[0]
    async with db_factory() as db:
        calls = list((await db.scalars(select(ResourceCall).where(ResourceCall.job_id == job['job_id']))).all())
        assert sum(c.actual for c in calls) == 300
        step = await db.get(Step, (job['job_id'], 'task'))
        assert jobs.unpack(step.checkpoint)['runtime']['written'] == ['prepared.txt']
    # The worker is a distinct claimant; no foreground continuation is needed.
    assert await worker.run_once()
    assert (await jobs.snapshot(user.id, job['job_id']))['status'] == 'completed'
    assert tool.execute.await_count == 2  # one preparation + one final write
    assert await jobs.publish_deliveries() == 1
    assert await jobs.publish_deliveries() == 0
    async with db_factory() as db:
        messages = list((await db.scalars(select(Message).where(Message.session_id == session.id))).all())
        assert sum(m.role == 'user' for m in messages) == 1
        usage = list((await db.scalars(select(UsageRecord).where(UsageRecord.session_id == session.id))).all())
        assert sum(u.prompt_tokens + u.completion_tokens for u in usage) == 445
        assert sum(u.counts_as_turn for u in usage) == 1


async def test_promoted_unknown_usage_keeps_reservation(stores, db_factory, tmp_path):
    from sqlalchemy import select
    from claw.jobs.models import Job, ResourceCall
    runtime, jobs, worker, user, session, provider = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content=None, tool_calls=[ToolCall('plan', 'update_plan', {
            'goal': 'Create a CSV file', 'steps': [
                {'step': 'Generate', 'status': 'pending'}, {'step': 'Verify', 'status': 'pending'}]})])],
    ])
    await runtime.handle_message(user.id, session.id, 'Proceed with the analysis', permission_mode='auto')
    snapshot = (await jobs.list_jobs(user.id, session.id))[0]
    async with db_factory() as db:
        job = await db.get(Job, snapshot['job_id'])
        calls = list((await db.scalars(select(ResourceCall).where(ResourceCall.job_id == job.id))).all())
        assert len(calls) == 1 and calls[0].actual is None
        assert job.reserved_tokens == calls[0].reserved > 0
        assert job.active_seconds > 0
    assert len(provider.calls) == 1


async def test_short_turn_stays_foreground(stores, db_factory, tmp_path):
    from claw.jobs.provider import foreground_accounting
    runtime, jobs, _, user, session, provider = await prepare(stores, db_factory, tmp_path,
        [[ChatResult(content='Hello', usage={'prompt_tokens': 10, 'completion_tokens': 2})]])
    assert await runtime.handle_message(user.id, session.id, 'hi') == 'Hello'
    assert not await jobs.list_jobs(user.id, session.id)
    assert foreground_accounting.get() is None


async def test_research_promotion_retains_fetched_sources(stores, db_factory, tmp_path, monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock
    fetch = AsyncMock(return_value=SimpleNamespace(status_code=200, headers={'content-type': 'text/plain'}, text='Evidence'))
    monkeypatch.setattr('httpx.AsyncClient.get', fetch)
    runtime, jobs, worker, user, session, _ = await prepare(stores, db_factory, tmp_path, [
        [ChatResult(content=None, tool_calls=[ToolCall('a', 'web_fetch', {'url': 'https://one.test/source'})])],
        [ChatResult(content=None, tool_calls=[ToolCall('plan', 'update_plan', {
            'goal': 'Research and compare sources', 'steps': [
                {'step': 'Fetch next source', 'status': 'pending'}, {'step': 'Compare', 'status': 'pending'}]})])],
        [ChatResult(content=None, tool_calls=[ToolCall('again', 'web_fetch', {'url': 'https://one.test/source'}),
                                            ToolCall('b', 'web_fetch', {'url': 'https://two.test/source'})])],
        [ChatResult(content='[one](https://one.test/source) and [two](https://two.test/source)')],
    ])
    await runtime.handle_message(user.id, session.id, 'Please examine this topic')
    assert fetch.await_count == 1
    assert await worker.run_once()
    job = (await jobs.list_jobs(user.id, session.id))[0]
    assert job['status'] == 'completed', job
    assert fetch.await_count == 2  # repeated request served from the adopted source cache
