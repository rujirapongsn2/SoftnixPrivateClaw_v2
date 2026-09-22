"""Queryable Knowledge ingestion, calculations, access control and upload rules."""

import json

import pytest
from openpyxl import Workbook

from claw.db.stores import GroupStore, KnowledgeStore, UserStore
from claw.knowledge.dataset import DatasetError, execute_dataset_query, ingest_dataset, run_query
from tests.conftest_app import build_api_app, client


async def _register(c, email: str):
    response = await c.post(
        "/api/auth/register", json={"email": email, "password": "password123"}
    )
    return response.json()["access_token"]


def _bearer(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def test_csv_ingest_and_exact_aggregate(tmp_path):
    source = tmp_path / "sales.csv"
    source.write_text(
        'sku,category,qty,sales\n001,Drink,2,30.50\n002,Snack,3,"1,045.00"\n003,Drink,1,10.00\n',
        encoding="utf-8",
    )
    destination = tmp_path / "dataset"

    schema, rows = ingest_dataset(source, source.name, destination)

    assert rows == 3
    assert schema["tables"][0]["name"] == "data"
    sku = next(column for column in schema["tables"][0]["columns"] if column["name"] == "sku")
    assert sku["inferred_type"] == "text"  # Preserve leading-zero product codes.

    result = run_query(
        destination / "dataset.duckdb",
        schema,
        {
            "table": "data",
            "group_by": ["category"],
            "aggregations": [
                {"function": "sum", "column": "sales", "alias": "total_sales"},
                {"function": "sum", "column": "qty", "alias": "total_qty"},
            ],
            "limit": 20,
        },
    )
    by_category = {row[0]: row[1:] for row in result["rows"]}
    assert by_category["Drink"] == ["40.500000", "3.000000"]
    assert by_category["Snack"] == ["1045.000000", "3.000000"]

    with pytest.raises(DatasetError, match="unknown column"):
        run_query(
            destination / "dataset.duckdb",
            schema,
            {"table": "data", "columns": ['sales" FROM data; DROP TABLE data; --']},
        )


def test_xlsx_ingest_creates_one_table_per_sheet(tmp_path):
    source = tmp_path / "inventory.xlsx"
    workbook = Workbook()
    active = workbook.active
    active.title = "Stock Bangkok"
    active.append(["SKU", "On hand"])
    active.append(["0007", 12])
    north = workbook.create_sheet("North")
    north.append(["SKU", "On hand"])
    north.append(["A-9", 4])
    workbook.save(source)

    schema, rows = ingest_dataset(source, source.name, tmp_path / "xlsx-dataset")

    assert rows == 2
    assert [table["name"] for table in schema["tables"]] == ["stock_bangkok", "north"]
    assert schema["tables"][0]["columns"][0]["inferred_type"] == "text"


def test_thai_headers_are_preserved_and_semicolon_csv_is_detected(tmp_path):
    source = tmp_path / "ยอดขาย.csv"
    source.write_text("รหัสสินค้า;ยอดขาย\n001;100\n", encoding="utf-8")

    schema, rows = ingest_dataset(source, source.name, tmp_path / "thai-dataset")

    assert rows == 1
    assert [column["name"] for column in schema["tables"][0]["columns"]] == [
        "รหัสสินค้า",
        "ยอดขาย",
    ]


def test_xlsx_formula_without_cached_value_is_rejected(tmp_path):
    source = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["qty", "price", "total"])
    sheet.append([2, 10, "=A2*B2"])
    workbook.save(source)

    with pytest.raises(DatasetError, match="formula cell"):
        ingest_dataset(source, source.name, tmp_path / "formula-dataset")


def test_mixed_numeric_data_cannot_produce_silent_totals_or_range_filters(tmp_path):
    source = tmp_path / "mixed.csv"
    source.write_text("sales\n10\nnot-available\n20\n", encoding="utf-8")
    destination = tmp_path / "mixed-dataset"
    schema, _ = ingest_dataset(source, source.name, destination)

    column = schema["tables"][0]["columns"][0]
    assert column == {
        "name": "sales",
        "inferred_type": "text",
        "non_empty_values": 3,
        "numeric_values": 2,
        "datetime_values": 0,
    }
    with pytest.raises(DatasetError, match="not consistently numeric"):
        run_query(
            destination / "dataset.duckdb",
            schema,
            {"table": "data", "aggregations": [{"function": "sum", "column": "sales"}]},
        )
    with pytest.raises(DatasetError, match="mixed value types"):
        run_query(
            destination / "dataset.duckdb",
            schema,
            {"table": "data", "aggregations": [{"function": "min", "column": "sales"}]},
        )
    with pytest.raises(DatasetError, match="range comparison"):
        run_query(
            destination / "dataset.duckdb",
            schema,
            {
                "table": "data",
                "columns": ["sales"],
                "filters": [{"column": "sales", "operator": "gt", "value": 9}],
            },
        )


def test_min_max_support_dates_and_text(tmp_path):
    source = tmp_path / "events.csv"
    source.write_text(
        "date,store\n2026-01-02T00:00:00,Bangkok\n2026-03-04T00:00:00,Chiang Mai\n",
        encoding="utf-8",
    )
    destination = tmp_path / "events-dataset"
    schema, _ = ingest_dataset(source, source.name, destination)

    result = run_query(
        destination / "dataset.duckdb",
        schema,
        {
            "table": "data",
            "aggregations": [
                {"function": "max", "column": "date", "alias": "latest"},
                {"function": "min", "column": "store", "alias": "first_store"},
            ],
        },
    )
    assert result["rows"] == [["2026-03-04 00:00:00", "Bangkok"]]


async def test_queryable_dataset_respects_group_sharing(db_factory, tmp_path):
    users = UserStore(db_factory)
    groups = GroupStore(db_factory)
    knowledge = KnowledgeStore(db_factory, is_postgres=False)
    sales = await groups.create("query-sales")
    other = await groups.create("query-other")
    owner = await users.create(email="query-owner@x.y", password_hash="h", group_id=sales.id)
    teammate = await users.create(email="query-mate@x.y", password_hash="h", group_id=sales.id)
    outsider = await users.create(email="query-outsider@x.y", password_hash="h", group_id=other.id)
    base = await knowledge.create_base(
        owner.id, "Sales Data", visibility="group", kind="queryable"
    )

    source = tmp_path / "sales.csv"
    source.write_text("region,sales\nBangkok,100\nNorth,50\n", encoding="utf-8")
    destination = tmp_path / base.id / "datasets" / "doc"
    schema, rows = ingest_dataset(source, source.name, destination)
    doc = await knowledge.create_pending_doc(
        kb_id=base.id,
        title="sales",
        filename=source.name,
        mime="text/csv",
        size=source.stat().st_size,
    )
    await knowledge.finalize_dataset_doc(
        doc_id=doc.id,
        dataset_path=str((destination / "dataset.duckdb").relative_to(tmp_path)),
        dataset_schema=schema,
        dataset_rows=rows,
    )

    teammate_result = json.loads(
        await execute_dataset_query(
            knowledge,
            tmp_path,
            teammate.id,
            knowledge_base="Sales Data",
            action="query",
            dataset="sales.csv",
            table="data",
            aggregations=[{"function": "sum", "column": "sales", "alias": "sales_total"}],
        )
    )
    outsider_result = json.loads(
        await execute_dataset_query(
            knowledge, tmp_path, outsider.id, knowledge_base="Sales Data", action="inspect"
        )
    )

    assert teammate_result["rows"] == [["150.000000"]]
    assert outsider_result == {"error": "Queryable Knowledge not found"}


async def test_api_enforces_knowledge_file_types(db_factory):
    app = build_api_app(db_factory)
    async with client(app) as c:
        token = await _register(c, "file-rules@x.io")
        headers = _bearer(token)
        general = (
            await c.post(
                "/api/knowledge",
                json={"name": "Docs", "kind": "general"},
                headers=headers,
            )
        ).json()
        queryable = (
            await c.post(
                "/api/knowledge",
                json={"name": "Data", "kind": "queryable"},
                headers=headers,
            )
        ).json()

        assert general["kind"] == "general"
        assert queryable["kind"] == "queryable"

        changed = await c.patch(
            f"/api/knowledge/{general['id']}",
            json={"kind": "queryable"},
            headers=headers,
        )
        assert changed.status_code == 200
        assert changed.json()["kind"] == "queryable"
        changed_back = await c.patch(
            f"/api/knowledge/{general['id']}",
            json={"kind": "general"},
            headers=headers,
        )
        assert changed_back.status_code == 200
        assert changed_back.json()["kind"] == "general"

        accepted_csv = await c.post(
            f"/api/knowledge/{general['id']}/documents",
            files={"files": ("notes.csv", b"name,value\nA,1\n", "text/csv")},
            headers=headers,
        )
        assert accepted_csv.status_code == 202
        locked_type = await c.patch(
            f"/api/knowledge/{general['id']}",
            json={"kind": "queryable"},
            headers=headers,
        )
        assert locked_type.status_code == 409
        assert "cannot be changed after files are uploaded" in locked_type.json()["detail"]

        general_xlsx = await c.post(
            f"/api/knowledge/{general['id']}/documents",
            files={"files": ("sales.xlsx", b"not-needed", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )
        queryable_pdf = await c.post(
            f"/api/knowledge/{queryable['id']}/documents",
            files={"files": ("notes.pdf", b"not-needed", "application/pdf")},
            headers=headers,
        )

        assert general_xlsx.status_code == 415
        assert "only in Queryable" in general_xlsx.json()["detail"]
        assert queryable_pdf.status_code == 415
        assert "CSV and XLSX" in queryable_pdf.json()["detail"]


async def test_api_upload_ingests_queryable_csv_end_to_end(db_factory):
    app = build_api_app(db_factory)
    state = app.state.claw
    await state.knowledge_service.start()
    try:
        async with client(app) as c:
            token = await _register(c, "dataset-upload@x.io")
            headers = _bearer(token)
            created = (
                await c.post(
                    "/api/knowledge",
                    json={"name": "Live Sales", "kind": "queryable"},
                    headers=headers,
                )
            ).json()
            upload = await c.post(
                f"/api/knowledge/{created['id']}/documents",
                files={"files": ("sales.csv", b"store,sales\nA,10\nB,20\n", "text/csv")},
                headers=headers,
            )
            assert upload.status_code == 202

            await state.knowledge_service._queue.join()
            docs = (
                await c.get(
                    f"/api/knowledge/{created['id']}/documents", headers=headers
                )
            ).json()

            assert docs[0]["status"] == "ready"
            assert docs[0]["dataset_rows"] == 2
            assert docs[0]["chunks"] == 0
    finally:
        await state.knowledge_service.stop()


async def test_document_delete_is_scoped_to_owned_knowledge_base(db_factory):
    app = build_api_app(db_factory)
    state = app.state.claw
    await state.knowledge_service.start()
    try:
        async with client(app) as c:
            victim_token = await _register(c, "delete-victim@x.io")
            attacker_token = await _register(c, "delete-attacker@x.io")
            victim_headers = _bearer(victim_token)
            attacker_headers = _bearer(attacker_token)
            victim_base = (
                await c.post(
                    "/api/knowledge",
                    json={"name": "Victim Data", "kind": "queryable", "visibility": "public"},
                    headers=victim_headers,
                )
            ).json()
            attacker_base = (
                await c.post(
                    "/api/knowledge",
                    json={"name": "Attacker Data", "kind": "queryable"},
                    headers=attacker_headers,
                )
            ).json()
            await c.post(
                f"/api/knowledge/{victim_base['id']}/documents",
                files={"files": ("sales.csv", b"store,sales\nA,10\n", "text/csv")},
                headers=victim_headers,
            )
            await state.knowledge_service._queue.join()
            victim_docs = (
                await c.get(
                    f"/api/knowledge/{victim_base['id']}/documents",
                    headers=attacker_headers,
                )
            ).json()
            victim_doc_id = victim_docs[0]["id"]

            response = await c.delete(
                f"/api/knowledge/{attacker_base['id']}/documents/{victim_doc_id}",
                headers=attacker_headers,
            )

            assert response.status_code == 404
            assert await state.knowledge.get_doc(victim_doc_id) is not None
    finally:
        await state.knowledge_service.stop()


async def test_dataset_is_removed_when_document_disappears_during_finalize(
    db_factory, monkeypatch
):
    app = build_api_app(db_factory)
    state = app.state.claw
    owner = await state.users.get_or_create_by_email("race-owner@x.io")
    base = await state.knowledge.create_base(owner.id, "Race Data", kind="queryable")
    staged = state.knowledge_service.staging_dir / "race-upload.part"
    staged.write_text("store,sales\nA,10\n", encoding="utf-8")

    async def delete_before_finalize(*, doc_id, **_):
        await state.knowledge.delete_doc(doc_id)
        return None

    monkeypatch.setattr(state.knowledge, "finalize_dataset_doc", delete_before_finalize)
    await state.knowledge_service.start()
    try:
        queued = await state.knowledge_service.enqueue_upload(
            base.id, "sales.csv", "text/csv", str(staged), staged.stat().st_size
        )
        await state.knowledge_service._queue.join()

        destination = state.knowledge_service.root / base.id / "datasets" / queued["id"]
        assert not destination.exists()
    finally:
        await state.knowledge_service.stop()
